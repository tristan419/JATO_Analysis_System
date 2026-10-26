"""Common-layout examples only; no customer or bank data in fixtures."""
import hashlib
import json
from copy import deepcopy
from io import BytesIO
from unittest.mock import Mock
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from app.core.security import UserContext
from app.core.security import get_current_user
from app.api.routes import order_genius_vehicle_allocation as route
from app.services import pi_invoice_templates as templates
from app.services import order_genius_vehicle_exporter as exporter


@pytest.fixture
def source(tmp_path, monkeypatch):
    monkeypatch.setattr(exporter, "PI_TEMPLATE_ROOT", tmp_path)
    wb = openpyxl.Workbook(); ws = wb.active
    for cell, value in {"A1": "Seller", "A2": "Seller address", "A3": "Buyer", "A4": "Buyer address",
                        "A5": "Bank", "A6": "CFR", "A7": "LC", "F1": "Date", "F2": "Reference"}.items():
        ws[cell] = value
    for col, label in enumerate(["No", "Model", "Version", "QTY.", "FOB Per Unit (EUR)", "Freight (EUR)", "Insurance (EUR)", "Total FOB", "TOTAL CFR (EUR)"], 1):
        ws.cell(10, col, label)
    ws["B11"] = "Example order"; ws["B12"] = "Old sample"; ws["A13"] = "Total"
    ws.merge_cells("A13:C13"); ws["A14"] = "Port of Shipment: Old"; ws["A15"] = "Port of Discharge: Old"
    ws.merge_cells("A14:I14"); ws.merge_cells("A15:I15")
    ws.column_dimensions["B"].width = 28
    ws.row_dimensions[11].height = 34
    output = BytesIO(); wb.save(output)
    content = output.getvalue()
    fields = {key: {"cell": cell, "value": value, "prefix": ""} for key, cell, value in [
        ("sellerName", "A1", "Seller"), ("sellerAddress", "A2", "Address"), ("buyerName", "A3", "Buyer"),
        ("buyerAddress", "A4", "Buyer address"), ("bankDetails", "A5", "Bank"), ("priceTerm", "A6", "CFR"),
        ("paymentTerm", "A7", "LC"), ("invoiceDate", "F1", ""), ("referenceNo", "F2", ""),
        ("portOfShipment", "A14", "Loading"), ("portOfDischarge", "A15", "Discharge")]}
    fields["portOfShipment"]["prefix"] = "Port of Shipment: "
    fields["portOfDischarge"]["prefix"] = "Port of Discharge: "
    analysis = templates.analyse(content)
    config = {key: analysis[key] for key in ("sheetName", "headerRow", "dataStart", "totalRow", "columns", "charges", "currency")}
    config.update(fields=fields, charges=["freightEur", "insuranceEur"])
    return content, dict(name="Country invoice", country="CZ", brands=["OMODA", "JAECOO"], shared=False,
                         sha256=hashlib.sha256(content).hexdigest(), config=config)


def detail():
    return {"header": {"piCode": "PI-CZ-202609-001", "countryCode": "CZ", "orderDate": "2026-09-01"},
            "lines": [{"piLineCode": "L01", "brand": "OMODA", "quantity": 2}],
            "vehicles": [{"piLineCode": "L01", "modelName": "MODEL", "version": "Trim", "powertrain": "BEV",
                          "materialCode": f"SAVED-{i}", "exteriorColorName": "White", "fobEur": 100.11,
                          "freightEur": 3.45, "insuranceEur": 0.67} for i in range(2)]}


def test_analysis_recognizes_table_and_never_saves(source):
    content, payload = source
    parsed = templates.analyse(content)
    assert (parsed["headerRow"], parsed["dataStart"], parsed["totalRow"]) == (10, 11, 13)
    assert parsed["columns"]["fobEur"] == "E"
    assert parsed["fields"]["portOfDischarge"]["cell"] == "A15"
    assert not templates.records()


@pytest.fixture
def split_label_invoice(source):
    """Real two-row/merged-label layout, with all commercial identities removed."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "PI-EXAMPLE"
    for address, value in {"A1": "Example sales company", "A3": "Example seller address", "B9": "M/S: Example buyer",
        "B10": "Address：Example buyer address", "I9": "Date:", "J9": "2026-10-09", "I10": "Ref. No:", "J10": "old-ref",
        "A21": "5-Price", "C21": "CIF Value EUR", "A22": "6- Payment", "C22": "Example payment conditions",
        "A24": "8- Bank Details:", "C24": "Example complete bank details", "A60": "Total",
        "A61": "10-Port of Shipment:Example loading", "A62": "11-Port of Discharge:Example discharge"}.items():
        ws[address] = value
    for address, label in {"B26": "Description of Goods", "F26": "QTY", "G26": "FOB Per Unit EUR", "H26": "Total FOB",
        "I26": "Freight EUR", "J26": "Insurance EUR", "K26": "Handing Charges EUR", "L26": "TOTAL CIF EUR",
        "B27": "Model", "C27": "Version", "D27": "Colour", "E27": "Material Code"}.items():
        ws[address] = label
    for merge in ("A1:L1", "A3:L3", "B9:E9", "B10:E10", "J9:L9", "J10:L10", "A21:B21", "C21:L21",
                  "A22:B22", "C22:L22", "A24:B24", "C24:L24", "B26:E26", "A59:K59", "A60:E60"):
        ws.merge_cells(merge)
    for col in "FGHIJKL": ws.merge_cells(f"{col}26:{col}27")
    for row in range(28, 59):
        for col, value in {"H": f"=G{row}*F{row}", "I": 1700, "J": 20, "K": 0,
                           "L": f"=(G{row}+I{row}+J{row}+K{row})*F{row}"}.items(): ws[f"{col}{row}"] = value
    for col in "FHL": ws[f"{col}60"] = f"=SUM({col}28:{col}58)"
    stream = BytesIO(); wb.save(stream)
    return stream.getvalue()


def test_split_labels_use_actual_value_cells_and_leave_unlabelled_identity_for_review(split_label_invoice):
    parsed = templates.analyse(split_label_invoice)
    assert (parsed["headerRow"], parsed["dataStart"], parsed["totalRow"]) == (26, 28, 60)
    assert {key: parsed["fields"][key]["cell"] for key in ("bankDetails", "priceTerm", "paymentTerm", "invoiceDate", "referenceNo")} == {
        "bankDetails": "C24", "priceTerm": "C21", "paymentTerm": "C22", "invoiceDate": "J9", "referenceNo": "J10"}
    assert parsed["fields"]["invoiceDate"]["prefix"] == ""
    assert parsed["fields"]["sellerName"]["cell"] == ""
    assert {"cell": "A1", "value": "Example sales company"} in parsed["sourceCells"]
    assert parsed["charges"] == []
    assert parsed["sourceCells"][-1] == {"cell": "L28", "value": "=(G28+I28+J28+K28)*F28"}
    with pytest.raises(ValueError, match="Seller / 卖方") as exc: templates.validate_config(split_label_invoice, parsed)
    assert "bankDetails" not in str(exc.value)
    assert not templates.records()


def test_confirmed_split_label_export_keeps_labels_and_uses_saved_pi_values(split_label_invoice):
    parsed = templates.analyse(split_label_invoice)
    parsed["fields"]["sellerName"] = {"cell": "A1", "value": "Example sales company", "prefix": ""}
    parsed["charges"] = ["freightEur", "insuranceEur", "handlingEur"]
    payload = dict(name="Example split label invoice", country="CZ", brands=["OMODA"], shared=False,
                   sha256=hashlib.sha256(split_label_invoice).hexdigest(), config=parsed)
    saved = templates.save(split_label_invoice, payload, "alice", False)
    template = templates.load_template(saved["id"], "alice", False, "CZ", {"OMODA"})
    pi = detail(); before = deepcopy(pi)
    out = exporter.generate_pi_invoice_excel(pi, {"invoiceDate": "2026-10-14"}, template)
    ws = openpyxl.load_workbook(out).active
    cached = openpyxl.load_workbook(BytesIO(out.getvalue()), data_only=True).active
    assert ws["I9"].value == "Date:" and ws["I10"].value == "Ref. No:"
    assert ws["J10"].value == pi["header"]["piCode"]
    assert str(ws["J9"].value).startswith("2026-10-14")
    assert ws["L28"].value == "=ROUND((G28+I28+J28+K28)*F28,2)"
    assert cached["L30"].value == 208.46 and cached["F30"].value == 2
    assert pi == before and template["path"].read_bytes() == split_label_invoice


def test_required_field_errors_are_aggregated_and_ambiguous_labels_are_not_guessed(source):
    content, payload = source
    config = deepcopy(payload["config"])
    for key in ("sellerName", "bankDetails", "priceTerm", "paymentTerm"): config["fields"][key]["value"] = ""
    with pytest.raises(ValueError) as exc: templates.validate_config(content, config)
    for label in ("卖方", "银行资料", "价格条款", "付款条款"): assert label in str(exc.value)
    wb = openpyxl.load_workbook(BytesIO(content)); ws = wb.active
    ws["A1"] = "Seller: One"; ws["A2"] = "Seller: Two"
    ws["F1"] = "Date:"; ws["G1"] = "Ref. No:"
    stream = BytesIO(); wb.save(stream)
    parsed = templates.analyse(stream.getvalue())
    assert parsed["fields"]["sellerName"]["cell"] == ""
    assert parsed["fields"]["invoiceDate"]["cell"] == ""


def test_personal_shared_admin_scope_and_independent_order(source):
    content, payload = source
    first = templates.save(content, payload, "alice", False)
    second = templates.save(content, {**payload, "name": "Shared", "shared": True}, "bob", False)
    assert templates.catalog("carol", False, "CZ", {"OMODA"})["items"] == [second]
    with pytest.raises(PermissionError): templates.load_template(first["id"], "bob", False, "CZ", {"OMODA"})
    assert templates.load_template(first["id"], "admin", True, "CZ", {"OMODA"})["owner"] == "alice"
    assert not templates.catalog("alice", False, "SE", {"OMODA"})["items"]
    assert not templates.catalog("alice", False, "CZ", {"CHERY"})["items"]
    templates.reorder("alice", False, "CZ", {"OMODA"}, [second["id"], first["id"]])
    assert templates.catalog("alice", False, "CZ", {"OMODA"})["items"][0]["id"] == second["id"]
    assert templates.catalog("bob", False, "CZ", {"OMODA"})["items"] == [second]
    with pytest.raises(PermissionError): templates.save(content, {**payload, "id": second["id"]}, "alice", False)
    with pytest.raises(PermissionError): templates.reorder("alice", False, "CZ", {"OMODA"}, [], "bob")
    templates.reorder("admin", True, "CZ", {"OMODA"}, [second["id"], first["id"]], "alice")
    with pytest.raises(ValueError): templates.reorder("alice", False, "CZ", {"OMODA"}, [first["id"]])


def test_custom_invoice_uses_saved_snapshots_cached_formulas_and_footer_mapping(source):
    content, payload = source
    record = templates.save(content, payload, "alice", False)
    template = templates.load_template(record["id"], "alice", False, "CZ", {"OMODA"})
    pi = detail(); before = deepcopy(pi)
    output = exporter.generate_pi_invoice_excel(pi, {}, template)
    ws = openpyxl.load_workbook(output).active
    cached = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    assert ws["B11"].value == "MODEL"
    assert str(next(merge for merge in ws.merged_cells.ranges if merge.min_col == 2)) == "B11:B12"
    assert ws["I11"].value == "=ROUND((E11+F11+G11)*D11,2)"
    assert cached["I13"].value == 208.46
    assert cached["H13"].value == 200.22
    assert cached["D13"].value == 2
    assert ws["D13"].number_format == "#,##0"
    assert ws["I13"].number_format == "#,##0.00"
    assert ws["A14"].value == "Port of Shipment: Loading"
    assert ws["F2"].value == "PI-CZ-202609-001"
    assert ws.column_dimensions["B"].width == 28
    assert pi == before and template["path"].read_bytes() == content


def test_custom_total_format_is_preserved(source):
    content, payload = source
    wb = openpyxl.load_workbook(BytesIO(content)); wb.active["I13"].number_format = '"EUR "#,##0.00'
    output = BytesIO(); wb.save(output); content = output.getvalue()
    payload["sha256"] = hashlib.sha256(content).hexdigest()
    record = templates.save(content, payload, "alice", False)
    template = templates.load_template(record["id"], "alice", False, "CZ", {"OMODA"})
    ws = openpyxl.load_workbook(exporter.generate_pi_invoice_excel(detail(), {}, template)).active
    assert ws["I13"].number_format == '"EUR "#,##0.00'


def test_total_label_stays_in_source_cell_with_an_unused_header_column(source):
    content, payload = source
    wb = openpyxl.load_workbook(BytesIO(content)); ws = wb.active
    ws["J1"] = "Fixed header note"
    ws["A13"] = " 合计 / Total "
    out = BytesIO(); wb.save(out); content = out.getvalue()
    payload["sha256"] = hashlib.sha256(content).hexdigest()
    saved = templates.save(content, payload, "alice", False)
    template = templates.load_template(saved["id"], "alice", False, "CZ", {"OMODA"})
    output = exporter.generate_pi_invoice_excel(detail(), {}, template)
    ws = openpyxl.load_workbook(output).active
    assert ws["A13"].value == " 合计 / Total " and ws["J13"].value is None
    assert str(next(merge for merge in ws.merged_cells.ranges if merge.min_row == 13)) == "A13:C13"
    assert ws["J1"].value == "Fixed header note"
    assert template["path"].read_bytes() == content


def test_large_currency_totals_widen_without_reformatting_or_changing_values(source):
    content, payload = source
    wb = openpyxl.load_workbook(BytesIO(content)); ws = wb.active
    ws["I13"].number_format = '"EUR "#,##0.00'
    ws.column_dimensions["I"].width = 20.6
    out = BytesIO(); wb.save(out); content = out.getvalue()
    payload["sha256"] = hashlib.sha256(content).hexdigest()
    saved = templates.save(content, payload, "alice", False)
    template = templates.load_template(saved["id"], "alice", False, "CZ", {"OMODA"})
    plan = {"country": "CZ", "month": "2026-10", "lines": [{**detail()["vehicles"][0], "brand": "OMODA", "quantity": 54000}]}
    output = exporter.generate_monthly_invoice_excel(plan, {"freightEur": 0, "insuranceEur": 0}, template)
    ws = openpyxl.load_workbook(output).active
    values = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    assert ws["I12"].number_format == '"EUR "#,##0.00' and ws.column_dimensions["I"].width >= 22
    assert values["I12"].value == 5405940 and values["D12"].value == 54000
    assert ws.column_dimensions["B"].width == 28 and ws["A12"].value == "Total"


def test_table_shrinks_footer_moves_and_zero_overrides_do_not_mutate_pi(source):
    content, payload = source
    record = templates.save(content, payload, "alice", False)
    template = templates.load_template(record["id"], "alice", False, "CZ", {"OMODA"})
    pi = detail(); pi["vehicles"][1]["materialCode"] = "SAVED-0"
    pi["vehicles"][0]["freightEur"] = pi["vehicles"][1]["freightEur"] = None
    with pytest.raises(ValueError): exporter.generate_pi_invoice_excel(pi, {}, template)
    out = exporter.generate_pi_invoice_excel(pi, {"freightEur": 0, "insuranceEur": 0, "referenceNo": "=BAD()"}, template)
    ws = openpyxl.load_workbook(out).active
    cached = openpyxl.load_workbook(BytesIO(out.getvalue()), data_only=True).active
    assert ws["A13"].value == "Port of Shipment: Loading"
    assert ws["A14"].value == "Port of Discharge: Discharge"
    assert cached["I12"].value == 200.22
    assert ws["F2"].data_type == "s"
    assert pi["vehicles"][0]["freightEur"] is None


@pytest.mark.parametrize("change", ["missing", "overlap", "charge", "bounds", "currency"])
def test_invalid_mapping_blocks_save(source, change):
    content, payload = source
    config = payload["config"]
    if change == "missing": config["fields"]["buyerName"]["value"] = ""
    if change == "overlap": config["columns"]["total"] = "D"
    if change == "charge": config["charges"] = ["unknown"]
    if change == "bounds": config["fields"]["bankDetails"]["cell"] = "A11"
    if change == "currency": config["currency"] = "CHF"
    with pytest.raises(ValueError): templates.save(content, payload, "alice", False)
    assert not templates.records()


def test_changed_source_and_damaged_top_template_never_fall_back(source, monkeypatch):
    content, payload = source
    first = templates.save(content, payload, "alice", False)
    second = templates.save(content, {**payload, "name": "Second"}, "alice", False)
    templates.reorder("alice", False, "CZ", {"OMODA"}, [first["id"], second["id"]])
    templates._path(first["id"]).joinpath("source.xlsx").write_bytes(b"damaged")
    monkeypatch.setattr(route, "_template_scope", Mock(return_value=(detail(), "CZ", {"OMODA"})))
    with pytest.raises(ValueError, match="damaged"):
        route._selected_template(Mock(), UserContext("order_filler", "alice"), "PI")


def test_scope_is_rechecked_before_any_template_write(source, monkeypatch):
    content, payload = source
    monkeypatch.setattr(route, "_management_scope", Mock(side_effect=HTTPException(403, "Revoked")))
    with pytest.raises(HTTPException): route._save_template(Mock(), UserContext("order_filler", "alice"), content, payload)
    assert not templates.records()


def test_country_library_can_save_before_any_pi_exists(source, monkeypatch):
    content, payload = source
    session = Mock()
    account = SimpleNamespace(primary_country_code="BG", secondary_country_codes=[], brands=["OMODA", "JAECOO"])
    session.query.return_value.filter.return_value.first.return_value = account
    user = UserContext("order_filler", "alice")
    monkeypatch.setattr(route, "_invoice_detail", Mock(side_effect=AssertionError("Must not read a PI")))
    assert route.list_scoped_invoice_templates("BG", "OMODA,JAECOO", "", session, user)["items"] == []
    saved = route._save_template(session, user, content, {**payload, "country": "BG"})
    assert route.list_scoped_invoice_templates("BG", "OMODA,JAECOO", "", session, user)["items"][0]["id"] == saved["id"]
    account.brands = []
    with pytest.raises(HTTPException) as exc:
        route._save_template(session, user, content, {**payload, "country": "BG"})
    assert exc.value.status_code == 403
    assert len(templates.records()) == 1


def test_monthly_http_export_reuses_uploaded_top_template_and_rechecks_scope(source, monkeypatch):
    content, payload = source
    saved = templates.save(content, {**payload, "name": "September reusable layout"}, "alice", False)
    session = Mock()
    account = SimpleNamespace(primary_country_code="CZ", secondary_country_codes=[], brands=["OMODA", "JAECOO"])
    session.query.return_value.filter.return_value.first.return_value = account
    plan = {"country": "CZ", "month": "2026-10", "lines": [{**detail()["vehicles"][0], "brand": "OMODA", "quantity": 54}]}
    monkeypatch.setattr(route, "build_monthly_invoice_plan", Mock(return_value=plan))
    monkeypatch.setattr(route, "_invoice_detail", Mock(side_effect=AssertionError("Monthly export must not read/create PI")))
    app = FastAPI(); app.include_router(route.router)
    app.dependency_overrides[route.get_db_session] = lambda: session
    app.dependency_overrides[get_current_user] = lambda: UserContext("order_filler", "alice")
    path = "/order-genius/vehicle-allocation/monthly-invoice?country=CZ&month=2026-10"
    with TestClient(app) as client:
        preview = client.get(path)
        assert preview.status_code == 200
        ctx = preview.json()["context"]
        assert ctx["templateKey"] == saved["id"] and ctx["unitCount"] == 54
        options = {**ctx["defaults"], "templateId": saved["id"], "templateRevision": ctx["templateRevision"],
                   "freightEur": 3.45, "insuranceEur": 0.67}
        response = client.post(path, json=options)
        assert response.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(response.content), data_only=True).active
        assert ws["D12"].value == 54 and ws["I12"].value == 5628.42
        assert ws["F2"].value == "PI-CZ-202610"
        assert ctx["freightModels"] == [{"brand": "OMODA", "modelName": "MODEL", "unitCount": 54}]
        grouped = {key: value for key, value in options.items() if key != "freightEur"}
        grouped["modelFreightEur"] = [{"brand": "OMODA", "modelName": "MODEL", "freightEur": 20}]
        model_response = client.post(path, json=grouped)
        assert model_response.status_code == 200
        model_ws = openpyxl.load_workbook(BytesIO(model_response.content), data_only=True).active
        assert model_ws["F11"].value == 20 and model_ws["I12"].value == 6522.12
        assert plan["lines"][0]["freightEur"] == 3.45
        plan["lines"].append({**plan["lines"][0], "modelName": "NEW MODEL", "quantity": 1})
        rejected = client.post(path, json=grouped)
        assert rejected.status_code == 409 and "NEW MODEL" in rejected.json()["detail"]
        plan["lines"].pop()
        assert templates.load_template(saved["id"], "alice", False, "CZ", {"OMODA"})["path"].read_bytes() == content
        plan.update(month="2026-11", lines=[{**plan["lines"][0], "quantity": 4}])
        later = client.get(path.replace("2026-10", "2026-11")).json()["context"]
        assert later["templateKey"] == saved["id"] and later["defaults"]["referenceNo"] == "PI-CZ-202611"
        assert later["unitCount"] == 4
        assert client.post(path, json={**options, "templateRevision": "stale"}).status_code == 409
        account.brands = []
        assert client.post(path, json=options).status_code == 403
        assert client.get(path).status_code == 403
    session.add.assert_not_called(); session.commit.assert_not_called()


@pytest.mark.parametrize("country,brands,status", [("", ["OMODA"], 400), ("BGR", ["OMODA"], 400),
    ("BG", [], 400), ("BG", ["unknown"], 400), ("BG", ["OMODA", "CHERY"], 403), ("CH", ["OMODA"], 403)])
def test_management_scope_rejects_empty_unknown_or_unassigned_targets(country, brands, status):
    session = Mock()
    session.query.return_value.filter.return_value.first.return_value = SimpleNamespace(
        primary_country_code="BG", secondary_country_codes=[], brands=["OMODA"])
    with pytest.raises(HTTPException) as exc:
        route._management_scope(session, UserContext("order_filler", "alice"), country, brands)
    assert exc.value.status_code == status


def test_scoped_http_upload_edit_copy_order_and_revoke_without_pi(source, monkeypatch):
    content, payload = source
    payload = {**payload, "country": "BG"}
    session = Mock()
    account = SimpleNamespace(primary_country_code="BG", secondary_country_codes=[], brands=["OMODA", "JAECOO"])
    session.query.return_value.filter.return_value.first.return_value = account
    app = FastAPI(); app.include_router(route.router)
    user = UserContext("order_filler", "alice")
    app.dependency_overrides[route.get_db_session] = lambda: session
    app.dependency_overrides[get_current_user] = lambda: user
    monkeypatch.setattr(route, "_invoice_detail", Mock(side_effect=AssertionError("No PI required")))
    path = "/order-genius/vehicle-allocation/invoice-templates"
    query = "?country=BG&brands=OMODA,JAECOO"
    with TestClient(app) as client:
        assert client.get(path + query).json()["items"] == []
        preview = client.post(path + "/preview" + query, files={"file": ("BG.xlsx", content)})
        assert preview.status_code == 200 and preview.json()["sha256"] == payload["sha256"]
        assert not templates.records()
        saved = client.post(path, data={"metadata": json.dumps(payload)}, files={"file": ("BG.xlsx", content)})
        assert saved.status_code == 200
        record = saved.json(); template_id = record["id"]
        assert client.patch(path + "/" + template_id, json={**payload, "name": "Updated"}).status_code == 200
        copy = client.post(path + "/" + template_id + "/copy" + query)
        assert copy.status_code == 200 and copy.json()["shared"] is False
        ids = [copy.json()["id"], template_id]
        ordered = client.put("/order-genius/vehicle-allocation/invoice-template-order", json={
            "country": "BG", "brands": payload["brands"], "ids": ids, "owner": "alice"})
        assert ordered.status_code == 200
        assert [item["id"] for item in client.get(path + query).json()["items"]] == ids
        assert client.get(path + query + "&owner=bob").status_code == 403
        assert client.get(path + "?country=BG&brands=OMODA,CHERY").status_code == 403
        account.brands = []
        assert client.post(path + "/preview" + query, files={"file": ("BG.xlsx", content)}).status_code == 403
        assert client.patch(path + "/" + template_id, json=payload).status_code == 403
        assert client.delete(path + "/" + template_id + query).status_code == 403
        assert len(templates.records()) == 2
        account.brands = payload["brands"]
        assert client.delete(path + "/" + template_id + query).status_code == 200
        user = UserContext("viewer", "alice")
        assert client.get(path + query).status_code == 403
        user = UserContext("developer", "admin")
        assert client.get(path + "?country=BG&brands=CHERY").status_code == 200


def test_no_macros_external_links_multiple_worksheets_or_bad_xlsx(source):
    content, _ = source
    with pytest.raises(ValueError): templates.analyse(b"not xlsx")
    wb = openpyxl.load_workbook(BytesIO(content)); wb.create_sheet("Other")
    out = BytesIO(); wb.save(out)
    with pytest.raises(ValueError, match="one invoice"): templates.analyse(out.getvalue())


def test_delete_only_owned_source_and_existing_copies_survive(source):
    content, payload = source
    original = templates.save(content, {**payload, "shared": True}, "alice", False)
    copy = templates.save(content, payload, "bob", False)
    with pytest.raises(PermissionError): templates.remove(original["id"], "bob", False, "CZ", {"OMODA"})
    templates._path(original["id"]).joinpath("source.xlsx").write_bytes(b"damaged")
    templates.remove(original["id"], "alice", False, "CZ", {"OMODA"})
    assert [record["id"] for record in templates.records()] == [copy["id"]]
    with pytest.raises(ValueError): templates.remove("../../", "admin", True, "CZ", {"OMODA"})


def test_export_rejects_changed_template_since_preview(source, monkeypatch):
    content, payload = source
    record = templates.save(content, payload, "alice", False)
    template = templates.load_template(record["id"], "alice", False, "CZ", {"OMODA"})
    monkeypatch.setattr(route, "_invoice_detail", Mock(return_value=detail()))
    monkeypatch.setattr(route, "_selected_template", Mock(return_value=template))
    with pytest.raises(HTTPException) as exc:
        route.export_pi_invoice("PI", {"templateId": record["id"], "templateRevision": "stale"}, Mock(), UserContext("order_filler", "alice"))
    assert exc.value.status_code == 409 and "changed since preview" in exc.value.detail
