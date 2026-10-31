from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import Column, MetaData, Table, create_engine, inspect, select
from sqlalchemy.dialects.postgresql import JSONB, UUID
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.routes import material_assets as route
from app.core.security import UserContext, get_current_user
from app.db.models import ConfigVersion, MaterialAssetBinding, MaterialAssetFile, MaterialAssetGroup, VehicleTrim
from app.db.session import get_db_session
from app.services import material_asset_service as svc

app = FastAPI()
app.include_router(route.router, prefix="/v1")


@compiles(JSONB, "sqlite")
def compile_jsonb(_type, _compiler, **_kw):
    return "JSON"


@pytest.fixture
def db():
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}, execution_options={"schema_translate_map": {"ordering": None, "engineering_config": None, "ops": None}})
    metadata = MetaData()
    Table("import_batches", metadata, Column("import_batch_id", UUID(as_uuid=True), primary_key=True), schema="ops")
    for cls in (VehicleTrim, ConfigVersion, MaterialAssetGroup, MaterialAssetFile, MaterialAssetBinding):
        cls.__table__.to_metadata(metadata)
    metadata.create_all(engine)
    with engine.connect() as conn:
        conn.exec_driver_sql("PRAGMA foreign_keys=ON")
        conn.connection.driver_connection.create_function("octet_length", 1, len)
    with Session(engine) as session:
        yield session
    engine.dispose()


def material(template="T7150D**WT1001", model="J7", code="CL"):
    return {"key": svc.material_key("JAECOO", model, template), "brand": "JAECOO", "modelName": model, "bomTemplate": template,
            "version": "Large battery FWD", "powertrain": "PHEV", "historical": False, "effectiveFrom": None, "effectiveTo": None,
            "colours": {f"code:{code}": {"key": f"code:{code}", "code": code, "name": "Black", "hex": "#000000"}},
            "interiors": {"name:BLACK-BLACK": {"key": "name:BLACK-BLACK", "code": None, "name": "Black-Black", "hex": None}}, "issues": [], "notes": []}


def group(db, model="J7", code="CL"):
    album = MaterialAssetGroup(brand="JAECOO", model_name=model, kind="exterior", colour_key=f"code:{code}", name="Black", created_by="editor")
    db.add(album); db.flush(); return album


def link(m, album, version=0):
    return {"materialKey": m["key"], "kind": "exterior", "colourKey": "code:CL", "groupId": str(album.asset_group_id), "versionId": None, "rowVersion": version}


def matrix(column=4, value="标配"):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = "在售可控资源表"
    sheet.cell(3, column, "JAECOO7-Exclusive"); sheet.cell(4, 2, "配置大类"); sheet.cell(4, 3, "配置子项")
    sheet.cell(5, 1, "F1"); sheet.cell(5, 2, "Safety"); sheet.cell(5, 3, "Airbag"); sheet.cell(5, column, value)
    out = BytesIO(); book.save(out); return out.getvalue()


def test_existing_parser_requires_explicit_column_and_retains_aa_provenance():
    original = matrix(27)
    assert svc.parse_configuration(original, "config.xlsx")["columns"][0]["column"] == 26
    parsed = svc.parse_configuration(original, "config.xlsx", 26)
    assert parsed["values"][0]["source_column"] == "AA"
    assert parsed["values"][0]["source_row"] == 5
    with pytest.raises(HTTPException) as error:
        svc.parse_configuration(original, "config.xlsx", 3)
    assert error.value.status_code == 422


@pytest.mark.parametrize("filename,content", [("config.pdf", b"pdf"), ("config.xlsx", b"not a workbook"), ("config.xlsx", b"")])
def test_bad_config_is_not_reported_saved(filename, content):
    with pytest.raises(HTTPException) as error:
        svc.parse_configuration(content, filename)
    assert error.value.status_code == 422


def test_batch_reuse_uses_same_model_and_saved_colour_not_hex(db):
    one, two = material(), material("T7150D**WT1002")
    materials = {m["key"]: m for m in [one, two]}; album = group(db)
    svc.replace_bindings(db, materials, [link(one, album), link(two, album)], "editor"); db.commit()
    assert len(db.scalars(select(MaterialAssetBinding)).all()) == 2
    assert len(db.scalars(select(MaterialAssetGroup)).all()) == 1
    other_model, other_colour = material("T7150D**WT1003", "J8"), material("T7150D**WT1004", code="BW")
    for wrong in [other_model, other_colour]:
        with pytest.raises(HTTPException) as error:
            svc.replace_bindings(db, {wrong["key"]: wrong, **materials}, [link(wrong, album)], "editor")
        assert error.value.status_code == 422


def test_stale_batch_leaves_every_binding_unchanged(db):
    one, two = material(), material("T7150D**WT1002"); mats = {m["key"]: m for m in [one, two]}; album = group(db)
    svc.replace_bindings(db, mats, [link(one, album), link(two, album)], "editor"); db.commit()
    changes = [dict(link(one, album, 1), groupId=None), dict(link(two, album, 0), groupId=None)]
    with pytest.raises(HTTPException) as error: svc.replace_bindings(db, mats, changes, "editor")
    assert error.value.status_code == 409
    db.rollback()
    assert all(b.asset_group_id == album.asset_group_id for b in db.scalars(select(MaterialAssetBinding)))


def test_files_are_deferred_and_unlink_does_not_delete_originals(db):
    m, album = material(), group(db)
    original = b"\x89PNG\r\n\x1a\nphoto"
    svc.add_photos(db, album, [("black.png", original)], "editor")
    svc.replace_bindings(db, {m["key"]: m}, [link(m, album)], "editor"); db.commit(); db.expire_all()
    file = db.scalars(select(MaterialAssetFile)).one(); assert "content" in inspect(file).unloaded
    svc.replace_bindings(db, {m["key"]: m}, [dict(link(m, album, 1), groupId=None)], "editor"); db.commit()
    assert file.content == original
    with pytest.raises(HTTPException): svc.group_access(db, album.asset_group_id, {m["key"]: m})
    assert svc.group_access(db, album.asset_group_id, {m["key"]: m}, editor=True) == album


def test_versions_are_immutable_and_missing_is_not_unequipped(db):
    m = material(); content = matrix()
    first = svc.create_configuration(db, m, svc.parse_configuration(content, "config.xlsx", 3), content, "v1.xlsx", "editor")
    second_content = matrix(value="-")
    second = svc.create_configuration(db, m, svc.parse_configuration(second_content, "config.xlsx", 3), second_content, "v2.xlsx", "editor")
    db.commit(); db.expire_all()
    assert second.version_no == 2 and second.parent_version_id == first.version_id
    assert first.values_snapshot_json[0]["raw_value"] == "标配"
    rows = svc.compare_versions([first, second, SimpleNamespace(values_snapshot_json=[])])
    assert rows[0]["values"][0]["raw_value"] == "标配" and rows[0]["values"][2] is None


def test_empty_country_scope_never_queries_unrestricted_bom(monkeypatch):
    monkeypatch.setattr(svc, "ordering_countries", lambda *_: set())
    monkeypatch.setattr(svc.bom_repo, "list_all_material_skus_for_admin", lambda *_args, **_kw: pytest.fail("Unrestricted call"))
    assert svc.catalog_materials(None, UserContext("order_filler", "limited")) == {}


def test_catalog_reads_saved_base_colours_and_rejects_conflicting_dates(monkeypatch):
    from datetime import date
    monkeypatch.setattr(svc, "ordering_countries", lambda *_: {"CH"})
    monkeypatch.setattr(svc, "ordering_brands", lambda *_: {"JAECOO"})
    def sku(start): return SimpleNamespace(bom_template="T7150D**WT1001", brand="JAECOO", model_name="J7", version="FWD", powertrain="PHEV", lifecycle_status="active", is_active=True, effective_from_date=start, effective_to_date=None, effective_from_month=None, effective_to_month=None, exterior_color_code="CL", exterior_color_name="Black", colour_hex="#000000", interior_colour_code=None, interior_color_name="Black-Black", remark="")
    monkeypatch.setattr(svc.bom_repo, "list_all_material_skus_for_admin", lambda *_a, **_kw: [sku(date(2026, 1, 1)), sku(date(2026, 2, 1))])
    result = next(iter(svc.catalog_materials(None, UserContext("editor", "e")).values()))
    assert result["issues"] and result["colours"]["code:CL"]["hex"] == "#000000"


def test_routes_enforce_roles_and_authorize_every_download(db, monkeypatch):
    m = material(); album = group(db); svc.add_photos(db, album, [("a.png", b"\x89PNG\r\n\x1a\nabc")], "e"); db.commit()
    monkeypatch.setattr(svc, "catalog_materials", lambda *_: {m["key"]: m})
    app.dependency_overrides[get_db_session] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: UserContext("order_filler", "f")
    try:
        client = TestClient(app)
        assert client.post("/v1/order-genius/material-assets/groups", data={}).status_code == 403
        assert client.get(f"/v1/order-genius/material-assets/files/{db.scalars(select(MaterialAssetFile)).one().asset_file_id}").status_code == 404
        svc.replace_bindings(db, {m["key"]: m}, [link(m, album)], "e"); db.commit()
        file = db.scalars(select(MaterialAssetFile)).one()
        assert client.get(f"/v1/order-genius/material-assets/files/{file.asset_file_id}").content.startswith(b"\x89PNG")
        assert client.get(f"/v1/order-genius/material-assets/download?fileIds={file.asset_file_id},{uuid4()}").status_code == 404
        app.dependency_overrides[get_current_user] = lambda: UserContext("viewer", "v")
        assert client.get("/v1/order-genius/material-assets").status_code == 403
    finally: app.dependency_overrides.clear()


def test_referenced_group_and_config_cannot_be_deleted(db, monkeypatch):
    m = material(); mats = {m["key"]: m}; album = group(db)
    svc.replace_bindings(db, mats, [link(m, album)], "e"); db.commit()
    monkeypatch.setattr(svc, "catalog_materials", lambda *_: mats)
    with pytest.raises(HTTPException) as error: route.delete_group(album.asset_group_id, route.DeleteRequest(rowVersion=1, confirmed=True), db, UserContext("editor", "e"))
    assert error.value.status_code == 409
    content = matrix(); version = svc.create_configuration(db, m, svc.parse_configuration(content, "x.xlsx", 3), content, "x.xlsx", "e")
    svc.replace_bindings(db, mats, [{"materialKey": m["key"], "kind": "config", "colourKey": "", "groupId": None, "versionId": str(version.version_id), "rowVersion": 0}], "e"); db.commit()
    with pytest.raises(HTTPException): route.delete_config(version.version_id, True, db, UserContext("editor", "e"))
