"""Real SQLAlchemy transactions: exercise public pricing paths, not mock prices."""
from datetime import date, timedelta
from uuid import uuid4
from types import SimpleNamespace

import pytest
from sqlalchemy import MetaData, create_engine, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session

from app.db import models
from app.infra import order_genius_repository as repo
from app.services import order_genius_service as service
from app.services import order_genius_vehicle_service as vehicle_service
from app.services import material_replacement_service as replacements
from fastapi import HTTPException
from app.api.routes import order_genius as routes
from app.services.bom_admin_export_service import build_bom_admin_workbook
from app.services.order_quantity_parser import OrderQuantityImport, ImportedRow, ImportedQuantityCell
import openpyxl


@compiles(JSONB, "sqlite")
def compile_jsonb_sqlite(_type, _compiler, **_kw):
    return "JSON"


@pytest.fixture
def db():
    engine = create_engine("sqlite://", execution_options={"schema_translate_map": {"ordering": None}})
    tables = [models.MaterialBaselineVersion, models.MaterialSkuMaster,
              models.CountrySkuFobResolved, models.FobResolvedHistory,
              models.BrandColourSurchargeRule, models.SpecialColourSurchargeRule,
              models.BrandColourSwatchRule, models.CountryPaymentTermMaster,
              models.CountryMaterialFinance, models.CountryFobSourceMapping,
              models.CountryTemplateFobPeriod, models.OrderQuantityCell,
              models.QuantityCellHistory, models.MaterialSkuRemarkHistory, models.MaterialLifecycle,
              models.MaterialReplacement, models.MaterialNoticeRead]
    metadata = MetaData()
    for model in tables:
        table = model.__table__.to_metadata(metadata)
        # SQLite supports partial indexes, but does not read postgresql_where.
        for index in table.indexes:
            predicate = index.dialect_options["postgresql"].get("where")
            if predicate is not None:
                index.dialect_options["sqlite"]["where"] = predicate
        table.create(engine)
    with Session(engine) as session:
        baseline = repo.create_baseline_version(session, "test", None, "test", "test")
        session.flush()
        session.info["baseline"] = baseline.baseline_version_id
        for brand, amount in [("OMODA", 200), ("JAECOO", 300)]:
            for tier in ("dual", "special"):
                session.add(models.BrandColourSurchargeRule(brand=brand, colour_type=tier, surcharge_eur=amount))
        session.commit()
        yield session
    engine.dispose()


def sku(db, code, tier, brand="JAECOO", template="T**001", model="JAECOO7 SHS", name="Colour"):
    value = models.MaterialSkuMaster(
        baseline_version_id=db.info["baseline"], material_code=code, brand=brand,
        model_name=model, version="Premium", powertrain="PHEV", bom_template=template,
        exterior_color_code=code, exterior_color_name=name, exterior_color_type=tier,
        colour_tier=tier, is_active=True, lifecycle_status="active",
    )
    db.add(value)
    db.flush()
    return value


def test_monthly_invoice_reads_plan_and_effective_month_price_without_any_db_writes(db):
    from sqlalchemy import event
    sku(db, "S", "single"); fob(db, "S", 1500, base=1500)
    service.update_quantity_cell(db, "CH", 2026, 10, "S", 54, "test", 1)
    db.commit()  # Finish fixture writes before observing the read-only export.
    statements = []
    def capture(_conn, _cursor, statement, _params, _context, _many):
        statements.append(statement.lstrip().upper())
    event.listen(db.bind, "before_cursor_execute", capture)
    try:
        # Internal planning read remains available; formal invoice routes require saved PIs.
        plan = service.build_monthly_invoice_plan(db, "CH", 2026, 10, allowed_brands=None)
        assert [(line["materialCode"], line["quantity"], line["fobEur"]) for line in plan["lines"]] == [("S", 54, 1500)]
    finally:
        event.remove(db.bind, "before_cursor_execute", capture)
    assert statements and all(statement.startswith("SELECT") for statement in statements)
    assert all("PI_ORDER" not in statement and "PI_VEHICLE" not in statement for statement in statements)
    cell = db.scalars(select(models.OrderQuantityCell)).one()
    assert cell.quantity == 54 and cell.row_version == 1
    assert float(db.scalars(select(models.CountrySkuFobResolved)).one().final_fob_eur) == 1500


def test_monthly_invoice_has_no_remaining_quantity_or_current_date_dependency(db):
    sku(db, "S", "single"); fob(db, "S", 1500, base=1500)
    release_period(db, ["S"], start=date(2026, 10, 15), end=date(2026, 10, 25))
    service.update_quantity_cell(db, "CH", 2026, 10, "S", 7, "test", 1)
    plan = service.build_monthly_invoice_plan(db, "CH", 2026, 10, allowed_brands={"JAECOO"})
    assert plan["lines"][0]["quantity"] == 7 and plan["lines"][0]["fobEur"] == 1000
    assert plan["lines"][0]["months"]["10"]["requiresOrderDate"] is True
    with pytest.raises(ValueError, match="No positive"):
        service.build_monthly_invoice_plan(db, "CH", 2026, 9, allowed_brands={"JAECOO"})
    with pytest.raises(ValueError, match="No positive"):
        service.build_monthly_invoice_plan(db, "CH", 2026, 10, allowed_brands={"OMODA"})


@pytest.mark.parametrize("failure", ["no_price", "unreleased", "paused", "conflict"])
def test_monthly_invoice_never_omits_unquotable_positive_plan_rows(db, failure):
    sku(db, "S", "single"); price = fob(db, "S", 1500, base=1500)
    service.update_quantity_cell(db, "CH", 2026, 10, "S", 7, "test", 1)
    if failure == "no_price":
        db.delete(price)
    elif failure == "unreleased":
        repo.save_default_colour_scope(db, "CH", "T**001", [], "test")
    else:
        for start, end, base in [(1, 15, 0 if failure == "paused" else 1000), (16, 31, 0 if failure == "paused" else 1200)]:
            db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template="T**001", valid_from=date(2026, 10, start),
                valid_to=date(2026, 10, end), base_fob_eur=base, status="active"))
    db.commit()
    with pytest.raises(ValueError, match="S:"):
        service.build_monthly_invoice_plan(db, "CH", 2026, 10, allowed_brands=None)


def fob(db, code, value, base=None, surcharge=None, term="TT", country="CH", source="template_base"):
    row = models.CountrySkuFobResolved(
        baseline_version_id=db.info["baseline"], material_code=code, country_code=country,
        payment_term_code=term, final_fob_eur=value, base_fob_eur=base,
        uploaded_fob_eur=base, colour_surcharge_eur=surcharge, fob_source_mode=source,
    )
    db.add(row)
    db.flush()
    return row


def test_list_loading_query_count_does_not_grow_with_materials(db):
    from sqlalchemy import event
    counts = []
    for target in (8, 80):
        for index in range(0 if target == 8 else 8, target):
            code = f"LOAD{index:03}"
            sku(db, code, ("single", "dual", "special")[index % 3], template=f"T**{index // 8:03}")
            fob(db, code, 1300, base=1000)
        db.commit()
        current = []
        for read in [
            lambda: service.build_matrix_batch(db, ["CH", "SE"], 2026, selection_date=date(2026, 10, 15)),
            lambda: repo.list_bom_with_fob(db, include_conflicts=True),
            lambda: service.build_material_timeline(db, "CH", allowed_brands=None, today=date(2026, 10, 15)),
        ]:
            statements = []
            def capture(_conn, _cursor, statement, *_args):
                statements.append(statement.lstrip().upper())
            event.listen(db.bind, "before_cursor_execute", capture)
            try:
                result = read()
            finally:
                event.remove(db.bind, "before_cursor_execute", capture)
            assert result and statements and all(statement.startswith("SELECT") for statement in statements)
            current.append(len(statements))
        counts.append(current)
    assert counts[0] == counts[1]
    assert counts[1][0] <= 20 and counts[1][1] <= 11 and counts[1][2] <= 8


@pytest.mark.parametrize("tier", ["single", "dual", "special", None])
def test_preloaded_rules_match_single_sku_and_refresh_after_edits(db, tier):
    original = sku(db, "RULE", tier or "single")
    original.colour_tier = tier or ""
    # Model+colour beats brand+colour, including explicit zero; inactive never wins.
    repo.upsert_special_colour_surcharge(db, "JAECOO", "RULE", 150, colour_tier="special")
    exact = repo.upsert_special_colour_surcharge(db, "JAECOO", "RULE", 0, model_name=original.model_name, colour_tier="special")
    db.commit()
    expected = repo.resolve_colour_surcharge_for_sku(db, original, repo.resolve_effective_colour_tier(original))
    assert repo.resolve_colour_surcharges_for_skus(db, [original])[original.material_sku_id] == expected
    exact.is_active = False
    for rule in repo.list_colour_surcharges(db):
        rule.surcharge_eur = 450
    db.commit()
    assert repo.resolve_colour_surcharges_for_skus(db, [original])[original.material_sku_id] == repo.resolve_colour_surcharge_for_sku(
        db, original, repo.resolve_effective_colour_tier(original))
    if tier == "special":
        assert repo.resolve_colour_surcharges_for_skus(db, [original])[original.material_sku_id]["amount"] == 150
    elif tier == "dual":
        assert repo.resolve_colour_surcharges_for_skus(db, [original])[original.material_sku_id]["amount"] == 450


def test_preloaded_colour_rules_keep_same_code_in_different_baselines_distinct(db):
    first = sku(db, "SAME", "single")
    baseline = repo.create_baseline_version(db, "next", None, "test", "test")
    db.flush()
    db.info["baseline"] = baseline.baseline_version_id
    second = sku(db, "SAME", "dual")
    db.commit()
    pricing = repo.resolve_colour_surcharges_for_skus(db, [first, second])
    assert len(pricing) == 2
    assert pricing[first.material_sku_id]["amount"] == 0
    assert pricing[second.material_sku_id]["amount"] == 300


@pytest.mark.parametrize("codes", [None, [], ["S"]])
def test_preloaded_periods_match_date_resolution_and_read_latest(db, codes):
    original = sku(db, "S", "single")
    fob(db, "S", 1500, base=1500)
    for start, end, base in [(1, 10, 1000), (15, 20, 0), (25, 31, 1200)]:
        db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template=original.bom_template,
            valid_from=date(2026, 10, start), valid_to=date(2026, 10, end), base_fob_eur=base,
            status="active", allowed_colour_codes=codes))
    db.commit()
    periods = repo.list_country_template_fob_periods(db, "CH")
    pricing = repo.resolve_colour_surcharges_for_skus(db, [original])[original.material_sku_id]
    for day in (1, 10, 11, 15, 21, 25, 31):
        when = date(2026, 10, day)
        assert service.resolve_date_effective_fob(db, "CH", original, when, periods=periods, colour_pricing=pricing) == service.resolve_date_effective_fob(db, "CH", original, when)
    first = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 10, 1))
    periods[0].base_fob_eur = 1400
    db.commit()
    second = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 10, 1))
    if codes == []:
        assert second["rows"][0]["fobEur"] is None
    else:
        assert first["rows"][0]["fobEur"] != second["rows"][0]["fobEur"]


def test_preloaded_overlapping_periods_preserve_conflict_and_country_scope(db):
    original = sku(db, "S", "single")
    for country in ("CH", "SE"):
        for start in (1, 10):
            db.add(models.CountryTemplateFobPeriod(country_code=country, bom_template=original.bom_template,
                valid_from=date(2026, 10, start), valid_to=None, base_fob_eur=1200, status="active"))
    db.commit()
    periods = repo.list_country_template_fob_periods(db, "CH")
    assert len(periods) == 2 and {row.country_code for row in periods} == {"CH"}
    when = date(2026, 10, 15)
    expected = service.resolve_date_effective_fob(db, "CH", original, when)
    assert "Overlapping" in expected[1]["reason"]
    assert service.resolve_date_effective_fob(db, "CH", original, when, periods=periods) == expected


def release_period(db, codes, start=date(2026, 10, 1), end=date(2026, 10, 31), country="CH", period=None):
    row = repo.save_country_template_fob_period(
        db, country_code=country, bom_template="T**001", valid_from=start, valid_to=end,
        base_fob_eur=1000, remark=None, changed_by="test", allowed_colour_codes=codes,
        period_id=period.country_template_fob_period_id if period else None,
        row_version=period.row_version if period else None,
    )
    db.commit()
    return row


@pytest.mark.parametrize("codes", [None, [], ["S"], ["S", "D"]])
def test_default_colour_scope_has_no_price_or_dates_and_does_not_create_a_phase(db, codes):
    single = sku(db, "S", "single")
    sku(db, "D", "dual")
    price = fob(db, "S", 1500, base=1500)
    repo.save_default_colour_scope(db, "CH", "T**001", codes, "test")
    db.commit()
    db.expire_all()
    assert repo.list_default_colour_scopes(db, "CH")["T**001"] == (sorted(codes) if codes is not None else None)
    assert repo.list_country_template_fob_periods(db, "CH", "T**001") == []
    default = db.scalars(select(models.CountryTemplateFobPeriod)).one()
    assert default.status == "default" and default.valid_from is None and default.valid_to is None and default.base_fob_eur is None
    assert float(price.final_fob_eur) == 1500
    assert repo.is_default_colour_released(db, "CH", single) is (codes is None or "S" in codes)
    new = sku(db, "NEW", "single")
    assert repo.is_default_colour_released(db, "CH", new) is (codes is None)
    assert repo.is_default_colour_released(db, "SE", new)


def test_default_colour_only_patch_does_not_reprice_and_price_only_preserves_colours(db):
    sku(db, "S", "single")
    price = fob(db, "S", 1500, base=1500)
    db.commit()
    user = SimpleNamespace(name="test", role="admin")
    body = {"bomTemplate": "T**001", "countryCode": "CH", "materialCodes": ["S"]}
    result = routes.patch_bom_template_fob({**body, "allowedColourCodes": []}, db, user)
    assert result["details"] == [] and result["updated"] == 0
    assert float(price.final_fob_eur) == 1500
    assert not db.scalars(select(models.FobResolvedHistory)).all()
    routes.patch_bom_template_fob({**body, "baseFobEur": 1600}, db, user)
    assert repo.list_default_colour_scopes(db, "CH")["T**001"] == []
    assert float(price.final_fob_eur) == 1600
    with pytest.raises(HTTPException):
        routes.patch_bom_template_fob({**body, "baseFobEur": 1700, "allowedColourCodes": ["FOREIGN"]}, db, user)
    assert float(price.final_fob_eur) == 1600
    assert repo.list_default_colour_scopes(db, "CH")["T**001"] == []


def test_template_editor_reads_complete_scope_before_an_explicit_default_price_save(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1500, base=1500)
    fob(db, "D", 1800, base=1500)
    db.commit()
    user = SimpleNamespace(name="test", role="admin")
    scope = routes.get_bom_template_fob_periods("T**001", "CH", db, user)
    assert scope["materialCodes"] == ["D", "S"]
    assert not db.scalars(select(models.FobResolvedHistory)).all()
    result = routes.patch_bom_template_fob({"bomTemplate": "T**001", "countryCode": "CH",
        "materialCodes": scope["materialCodes"], "baseFobEur": 1600}, db, user)
    assert {row["materialCode"]: row["finalFobEur"] for row in result["details"]} == {"S": 1600, "D": 1900}


def test_default_colour_blocks_plan_import_pi_export_but_phase_is_independent(db):
    single = sku(db, "S", "single")
    price = fob(db, "S", 1500, base=1500)
    db.commit()
    parsed = OrderQuantityImport("CH", 2026, [ImportedRow("S", 99999, "Model", "Version", "Colour", [ImportedQuantityCell("S", 10, 4)])])
    assert service.apply_order_quantity_import(db, parsed, "test")["appliedCells"] == 1
    repo.save_default_colour_scope(db, "CH", "T**001", [], "test")
    db.commit()
    assert service.resolve_date_effective_fob(db, "CH", single, date(2026, 10, 15))[2]["status"] == "colour_not_released"
    quote = service.resolve_month_effective_fob(db, "CH", single, 2026, 10, default_fob=price, periods=[], uses_periods=False,
        colour_pricing={"amount": 0})
    assert not quote["isEditable"] and quote["status"] == "colour_not_released"
    assert not service.preview_order_quantity_import(db, parsed)["matchedRows"]
    assert service.apply_order_quantity_import(db, parsed, "test")["appliedCells"] == 0
    with pytest.raises(ValueError):
        service.update_quantity_cell(db, "CH", 2026, 10, "S", 5, "test", 1)
    with pytest.raises(HTTPException):
        vehicle_service._line_payload_from_material(db, "CH", "S", {}, pricing_date=date(2026, 10, 15))
    items, countries = repo.list_bom_with_fob(db)
    projected = service.prepare_bom_export_rows(db, items, countries, reference_date=date(2026, 10, 15))
    assert all(not row["exportPricesByCountry"] for row in projected)
    phase = release_period(db, ["S"])
    assert service.resolve_date_effective_fob(db, "CH", single, date(2026, 10, 15))[0] == 1000
    assert vehicle_service._line_payload_from_material(db, "CH", "S", {}, pricing_date=date(2026, 10, 15))["fobEur"] == 1000
    price.final_fob_eur = 0
    db.commit()
    quote = service.resolve_month_effective_fob(db, "CH", single, 2026, 10, default_fob=price, periods=[phase], uses_periods=True,
        colour_pricing={"amount": 0})
    assert quote["isEditable"] and quote["fobEur"] == 1000
    assert service.resolve_date_effective_fob(db, "CH", single, date(2026, 11, 15))[2]["status"] == "no_price"


def test_last_phase_deletion_previews_and_rechecks_saved_global_colours(db):
    sku(db, "S", "single")
    fob(db, "S", 1500, base=1500)
    repo.save_default_colour_scope(db, "CH", "T**001", ["S"], "test")
    phase = release_period(db, [])
    preview = repo.preview_country_template_fob_period_deletion(db, phase.country_template_fob_period_id, phase.row_version)
    assert preview["defaultAllowedColourCodes"] == ["S"]
    repo.save_default_colour_scope(db, "CH", "T**001", [], "test")
    db.commit()
    new = repo.preview_country_template_fob_period_deletion(db, phase.country_template_fob_period_id, phase.row_version)
    assert new["fingerprint"] != preview["fingerprint"]
    assert new["defaultAllowedColourCodes"] == []


@pytest.mark.parametrize("codes", [None, [], ["S"], ["S", "D"]])
def test_release_scope_persists_null_empty_and_explicit_lists(db, codes):
    single, dual = sku(db, "S", "single"), sku(db, "D", "dual")
    period = release_period(db, codes)
    period_id = period.country_template_fob_period_id
    db.expire_all()
    reloaded = db.get(models.CountryTemplateFobPeriod, period_id)
    expected = sorted(codes) if codes is not None else None
    assert reloaded.allowed_colour_codes == expected
    assert repo.country_template_fob_period_payload(reloaded)["allowedColourCodes"] == expected
    new = sku(db, "NEW", "single")
    assert repo.is_colour_released(reloaded, new) is (codes is None)
    assert repo.is_colour_released(reloaded, single) is (codes is None or "S" in codes)
    assert repo.is_colour_released(reloaded, dual) is (codes is None or "D" in codes)


@pytest.mark.parametrize("codes", ["S", [""], [12], ["OTHER-TEMPLATE"]])
def test_release_scope_rejects_bad_or_foreign_colour_codes(db, codes):
    sku(db, "S", "single")
    with pytest.raises(ValueError):
        release_period(db, codes)
    assert not db.scalars(select(models.CountryTemplateFobPeriod)).all()


def test_release_month_window_pi_date_and_other_country_share_same_rules(db):
    single = sku(db, "S", "single")
    fob(db, "S", 1000, base=1000)
    fob(db, "S", 1000, base=1000, country="SE")
    release_period(db, [], end=date(2026, 10, 14))
    release_period(db, ["S"], start=date(2026, 10, 15))
    quote = service.resolve_month_effective_fob(
        db, "CH", single, 2026, 10, default_fob=None,
        periods=repo.list_country_template_fob_periods(db, "CH", "T**001"), uses_periods=True,
        colour_pricing=repo.resolve_colour_surcharge_for_sku(db, single, "single"),
    )
    assert quote["isEditable"] and quote["requiresOrderDate"]
    assert quote["availableRanges"] == [{"from": "2026-10-15", "to": "2026-10-31"}]
    assert service.resolve_date_effective_fob(db, "CH", single, date(2026, 10, 14))[0] is None
    assert service.resolve_date_effective_fob(db, "CH", single, date(2026, 10, 15))[0] == 1000
    assert vehicle_service._line_payload_from_material(db, "SE", "S", {}, pricing_date=date(2026, 10, 14))["fobEur"] == 1000
    with pytest.raises(HTTPException, match="not released"):
        vehicle_service._line_payload_from_material(db, "CH", "S", {}, pricing_date=date(2026, 10, 14))
    assert vehicle_service._line_payload_from_material(db, "CH", "S", {}, pricing_date=date(2026, 10, 15))["fobEur"] == 1000
    with pytest.raises(HTTPException):
        vehicle_service._resolve_pi_pricing_date(db, "CH", "S", 2026, 10, None)


def test_import_preview_and_apply_recheck_release_without_changing_existing_plan(db):
    sku(db, "S", "single")
    fob(db, "S", 1000, base=1000)
    period = release_period(db, ["S"])
    parsed = OrderQuantityImport("CH", 2026, [
        ImportedRow("S", 99999, "Model", "Version", "Colour", [ImportedQuantityCell("S", 10, 4)]),
    ])
    preview = service.preview_order_quantity_import(db, parsed)
    assert not preview["matchedRows"][0]["cells"][0]["error"]
    assert service.apply_order_quantity_import(db, parsed, "test")["appliedCells"] == 1
    cell = repo.get_quantity_cell(db, "CH", 2026, 10, "S")
    assert cell.quantity == 4 and float(cell.fob_eur) == 1000
    parsed.rows[0].cells[0].quantity = 5
    assert not service.preview_order_quantity_import(db, parsed)["matchedRows"][0]["cells"][0]["error"]
    release_period(db, [], period=period)
    preview = service.preview_order_quantity_import(db, parsed)
    assert "未释放" in preview["matchedRows"][0]["cells"][0]["error"]
    rejected = service.apply_order_quantity_import(db, parsed, "test")
    assert rejected["appliedCells"] == 0 and rejected["skippedCells"] == 1
    assert cell.quantity == 4 and float(cell.fob_eur) == 1000
    with pytest.raises(ValueError, match="未释放"):
        service.update_quantity_cell(db, "CH", 2026, 10, "S", 5, "test", cell.row_version)


def test_export_filters_unpriced_unreleased_stopped_and_strikes_both_historical_prices(db):
    for code in ("S", "H", "NO-PRICE", "BLOCKED", "STOPPED"):
        sku(db, code, "single", template="T**001" if code in ("S", "BLOCKED") else f"{code}**001")
    for code in ("S", "H", "BLOCKED", "STOPPED"):
        fob(db, code, 1200, base=1200)
    release_period(db, ["S"])
    repo.save_country_template_fob_period(
        db, country_code="CH", bom_template="H**001", valid_from=date(2026, 10, 1),
        valid_to=date(2026, 10, 10), base_fob_eur=1500, remark=None, changed_by="test",
    )
    historical = repo.get_sku_by_material_code(db, "H")
    historical.effective_from_date, historical.effective_to_date = date(2026, 1, 1), date(2026, 10, 10)
    repo.save_country_template_fob_period(
        db, country_code="CH", bom_template="STOPPED**001", valid_from=date(2026, 10, 1),
        valid_to=None, base_fob_eur=0, remark=None, changed_by="test",
    )
    db.commit()
    items, countries = repo.list_bom_with_fob(db)
    projected = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 15))
    wb = openpyxl.load_workbook(build_bom_admin_workbook(projected, countries, country_code="CH", reference_date=date(2026, 10, 15)))
    ws = wb["JAECOO7 SHS"]
    rows = {ws.cell(r, 6).value: r for r in range(3, ws.max_row + 1) if ws.cell(r, 6).value}
    assert set(rows) == {"S", "H"}
    assert ws.cell(rows["S"], 12).value == 1200 and ws.cell(rows["S"], 13).value == 1000
    assert not ws.cell(rows["S"], 12).font.strike
    assert ws.cell(rows["H"], 12).value == 1200 and ws.cell(rows["H"], 13).value == 1500
    assert ws.cell(rows["H"], 12).font.strike and ws.cell(rows["H"], 13).font.strike
    assert "2026-10-10" in ws.cell(rows["H"], 11).value
    assert "2026-10-15" in ws["K1"].value
    assert all(r[17] == "historical" for r in wb["_BOM_ADMIN_DATA"].iter_rows(min_row=2, values_only=True) if r[6] == "H")
    assert wb["_BOM_ADMIN_SCHEMA"]["B8"].value == "2026-10-15"


def test_export_all_countries_keeps_one_country_without_exposing_unreleased_old_price(db):
    sku(db, "S", "single")
    fob(db, "S", 1000, base=1000, country="CH")
    fob(db, "S", 1200, base=1200, country="SE")
    release_period(db, [])
    items, countries = repo.list_bom_with_fob(db)
    projected = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 15))
    wb = openpyxl.load_workbook(build_bom_admin_workbook(projected, countries, reference_date=date(2026, 10, 15)))
    ws = wb["JAECOO7 SHS"]
    assert ws["L3"].value is None and ws["M3"].value is None
    assert ws["N3"].value == 1200 and ws["O3"].value == 1200
    assert wb["_BOM_ADMIN_DATA"].max_row == 2
    assert wb["_BOM_ADMIN_DATA"]["N2"].value == "SE"


@pytest.mark.parametrize("codes,base", [([], 1000), (None, 0)])
def test_historical_reference_never_bypasses_month_colour_release_or_stop(db, codes, base):
    historical = sku(db, "H", "single")
    fob(db, "H", 1200, base=1200)
    historical.effective_to_date = date(2026, 10, 10)
    repo.save_country_template_fob_period(
        db, country_code="CH", bom_template="T**001", valid_from=date(2026, 10, 1),
        valid_to=date(2026, 10, 10), base_fob_eur=base, allowed_colour_codes=codes,
        remark=None, changed_by="test",
    )
    db.commit()
    items, countries = repo.list_bom_with_fob(db)
    rows = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 15))
    assert rows[0]["exportPricesByCountry"] == {}
    with pytest.raises(ValueError, match="无可导出"):
        build_bom_admin_workbook(rows, countries, reference_date=date(2026, 10, 15))


def test_export_uses_authoritative_period_without_choosing_conflicting_saved_default(db):
    sku(db, "S", "single")
    fob(db, "S", 1200, base=1200, term="TT")
    fob(db, "S", 1300, base=1300, term="LC")
    release_period(db, ["S"])
    items, countries = repo.list_bom_with_fob(db)
    rows = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 15))
    price = rows[0]["exportPricesByCountry"]["CH"]
    assert price["savedFobEur"] is None and price["monthFobEur"] == 1000
    wb = openpyxl.load_workbook(build_bom_admin_workbook(rows, countries, reference_date=date(2026, 10, 15)))
    assert wb["JAECOO7 SHS"]["L3"].value is None and wb["JAECOO7 SHS"]["M3"].value == 1000


def test_explicit_base_links_accept_changed_trim_and_codes_within_model_without_changing_business_data(db):
    old = sku(db, "OLD-BW", "single", template="OLD**A", model="Old model")
    new = sku(db, "NEW-CL", "single", template="NEW**B", model="Old model")
    new.version = "New trim"
    old.effective_from_date, old.effective_to_date = date(2026, 1, 1), date(2026, 9, 30)
    fob(db, old.material_code, 1000, base=1000)
    db.commit()
    before = (old.material_code, old.row_version, old.effective_to_date, new.material_code)
    result = replacements.save_replacement(db, brand="JAECOO", source="OLD**A", target="NEW**B",
        note="New hardware", expected_version=0, actor="admin", allowed_brands=None)
    db.commit()
    assert result == {"changed": True, "rowVersion": 1}
    assert before == (old.material_code, old.row_version, old.effective_to_date, new.material_code)
    projection = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"}, today=date(2026, 10, 12))
    assert projection["links"] == [{"from": "OLD**A", "to": "NEW**B", "note": "New hardware"}]
    assert next(node for node in projection["nodes"] if node["bomTemplate"] == "NEW**B")["materialCodes"] == []
    assert service.build_material_timeline(db, "SE", allowed_brands=None)["nodes"] == []
    assert service.build_material_timeline(db, "CH", allowed_brands={"OMODA"})["nodes"] == []
    assert replacements.save_replacement(db, brand="JAECOO", source="OLD**A", target="NEW**B",
        note="New hardware", expected_version=1, actor="admin", allowed_brands=None)["changed"] is False
    assert repo.list_fobs_for_country_material_codes(db, "CH", [old.material_code])[old.material_code].final_fob_eur == 1000


def test_replacement_catalog_is_model_scoped_and_old_external_links_are_preserved(db):
    sku(db, "A", "single", template="A", model="Model one")
    sku(db, "B", "single", template="B", model="Model one")
    sku(db, "C", "single", template="C", model="Model two")
    existing = models.MaterialReplacement(brand="JAECOO", source_template="A", target_template="C", row_version=1,
                                         changed_by="legacy", changed_at=date(2026, 10, 1))
    db.add(existing); db.commit()
    catalog = replacements.replacement_catalog(db, {"JAECOO"}, "A")
    assert {node["bomTemplate"] for node in catalog["nodes"]} == {"A", "B"}
    assert catalog["links"][0]["to"] == "C"
    with pytest.raises(HTTPException, match="within this model"):
        replacements.save_replacement(db, brand="JAECOO", source="B", target="C", note="",
                                      expected_version=0, actor="admin", allowed_brands=None)
    assert existing.target_template == "C" and existing.row_version == 1


def test_month_export_includes_midmonth_release_even_if_month_end_is_stopped(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1200, base=1200)
    release_period(db, ["S"], start=date(2026, 10, 1), end=date(2026, 10, 10))
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 10, 11), valid_to=date(2026, 10, 31), base_fob_eur=0,
        allowed_colour_codes=[], remark=None, changed_by="test")
    db.commit()
    items, countries = repo.list_bom_with_fob(db)
    day = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 31))
    assert all(not row["exportPricesByCountry"] for row in day)
    month = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 31), whole_month=True)
    wb = openpyxl.load_workbook(build_bom_admin_workbook(month, countries, country_code="CH", reference_date=date(2026, 10, 31)))
    ws = wb["JAECOO7 SHS"]
    assert ws["F3"].value == "S" and ws["M3"].value == 1000
    assert ws["N3"].value == "2026-10-01 → 2026-10-10"
    assert "Month-end status" in ws["K1"].value
    assert ws.row_dimensions[1].height + ws.row_dimensions[2].height >= 80
    assert not any(row[5] == "D" for row in ws.iter_rows(min_row=3, values_only=True))


def test_month_export_zero_saved_historical_price_is_not_a_reference(db):
    value = sku(db, "S", "single")
    value.effective_to_date = date(2026, 9, 1)
    fob(db, "S", 0, base=0)
    db.commit()
    items, countries = repo.list_bom_with_fob(db)
    rows = service.prepare_bom_export_rows(db, items, countries, date(2026, 10, 31), whole_month=True)
    assert rows[0]["exportPricesByCountry"] == {}


def test_replacement_rejects_cycles_missing_nodes_cross_brand_and_duplicate_predecessor(db):
    for code in "ABC": sku(db, code, "single", template=code)
    sku(db, "OTHER", "single", template="OTHER", brand="OMODA")
    db.commit()
    def save(source, target, allowed=None):
        return replacements.save_replacement(db, brand="JAECOO", source=source, target=target,
            note="", expected_version=0, actor="admin", allowed_brands=allowed)
    save("A", "B"); db.commit()
    for source, target, expected in [("B", "A", "cycle"), ("C", "B", "predecessor"),
                                      ("A", "A", "itself"), ("C", "OTHER", "missing"), ("C", "ABSENT", "missing")]:
        with pytest.raises(HTTPException, match=expected): save(source, target)
    with pytest.raises(HTTPException) as denied: save("C", None, {"OMODA"})
    assert denied.value.status_code == 403


def test_personal_material_notices_are_scope_checked_and_account_local(db):
    value = sku(db, "A", "single", template="A")
    value.effective_from_date = date(2026, 10, 30)
    fob(db, "A", 1000, base=1000); db.commit()
    projection = service.build_material_timeline(db, "CH", allowed_brands=None, today=date(2026, 10, 12))
    alice, bob = uuid4(), uuid4()
    notices = replacements.material_notices(db, projection, user_id=alice, username="alice")
    assert len(notices) == 1 and notices[0]["read"] is False
    key = notices[0]["key"]
    replacements.acknowledge_material_notice(db, key=key, user_id=alice, notices=notices, read=True); db.commit()
    assert replacements.material_notices(db, projection, user_id=alice, username="alice")[0]["read"]
    assert not replacements.material_notices(db, projection, user_id=bob, username="bob")[0]["read"]
    denied = service.build_material_timeline(db, "SE", allowed_brands=None, today=date(2026, 10, 12))
    with pytest.raises(HTTPException):
        replacements.acknowledge_material_notice(db, key=key, user_id=alice, notices=replacements.material_notices(db, denied, user_id=alice, username="alice"), read=True)
    value.effective_from_date = date(2026, 10, 31); db.commit()
    updated = service.build_material_timeline(db, "CH", allowed_brands=None, today=date(2026, 10, 12))
    assert not replacements.material_notices(db, updated, user_id=alice, username="alice")[0]["read"]


def test_relationship_notice_excludes_actor_and_noop_does_not_renew_notice(db):
    sku(db, "A", "single", template="A"); sku(db, "B", "single", template="B")
    fob(db, "A", 1000, base=1000); db.commit()
    replacements.save_replacement(db, brand="JAECOO", source="A", target="B", note="Update",
        expected_version=0, actor="admin", allowed_brands=None); db.commit()
    projection = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert replacements.material_notices(db, projection, user_id=uuid4(), username="admin") == []
    notices = replacements.material_notices(db, projection, user_id=uuid4(), username="filler")
    assert len(notices) == 1 and notices[0]["target"] == "B"


def test_material_colour_tier_has_no_implicit_default() -> None:
    column = models.MaterialSkuMaster.__table__.c.colour_tier

    assert column.nullable is False
    assert column.default is None
    assert column.server_default is None


def test_product_save_is_atomic_and_returns_saved_powertrain(db):
    first = sku(db, "A", "single")
    second = sku(db, "B", "dual")
    first.model_name = second.model_name = "JAECOO8 SHS"
    fob(db, "A", 1000, base=1000)
    fob(db, "B", 1300, base=1000, surcharge=300)
    db.commit()
    body = {"materialCodes": ["A", "B"], "version": "7 seats", "powertrain": "HEV",
            "remark": "new note", "rowVersions": {"A": 1, "B": 99}}
    with pytest.raises(HTTPException) as error:
        routes.patch_sku_metadata("A", body, db, SimpleNamespace(name="test", role="admin"))
    assert error.value.status_code == 409
    db.expire_all()
    assert first.remark is None and first.row_version == 1
    assert first.version == "Premium" and first.powertrain == "PHEV"
    body["rowVersions"]["B"] = 1
    result = routes.patch_sku_metadata("A", body, db, SimpleNamespace(name="test", role="admin"))
    db.expire_all()
    assert result["productFields"]["powertrain"] == "HEV"
    assert {row["powertrain"] for row in service.build_matrix(db, "CH", 2026)["rows"]} == {"HEV"}
    assert first.powertrain == second.powertrain == "HEV"
    assert first.remark == second.remark == "new note"
    assert first.row_version == second.row_version == 2
    assert len(db.scalars(select(models.MaterialSkuRemarkHistory)).all()) == 2
    # Re-saving unchanged remarks neither fabricates a version nor adds history.
    routes.patch_sku_metadata("A", body, db, SimpleNamespace(name="test", role="admin"))
    db.expire_all()
    assert first.row_version == second.row_version == 2
    assert len(db.scalars(select(models.MaterialSkuRemarkHistory)).all()) == 2


@pytest.mark.parametrize("saved,expected", [("HEV", "HEV"), ("Other", "OTHER"), (None, "")])
def test_powertrain_survives_name_template_and_lifecycle_edits(db, saved, expected):
    original = sku(db, "ORIGINAL", "single", template="T**0008", model="OMODA5 HEV")
    original.powertrain = saved
    fob(db, original.material_code, 1000, base=1000)
    db.commit()
    # The same metadata path is used when editing an upgraded template.
    repo.update_sku_metadata(db, [original.material_code], model_name="OMODA5 BEV", version="Upgraded")
    original.bom_template = "T**0011"
    original.lifecycle_status = "historical"
    db.commit()
    db.expire_all()
    rows, _ = repo.list_bom_with_fob(db)
    assert rows[0]["powertrain"] == expected
    assert original.powertrain == saved
    assert repo.get_sku_by_material_code(db, original.material_code).powertrain == saved


@pytest.mark.parametrize("source_tier", ["single", "dual", "special"])
@pytest.mark.parametrize("target_tier", ["single", "dual", "special"])
@pytest.mark.parametrize("brand,amount", [("OMODA", 200), ("JAECOO", 300)])
def test_create_every_tier_from_every_source(db, source_tier, target_tier, brand, amount):
    source = sku(db, "SOURCE", source_tier, brand=brand)
    target = sku(db, "NEW", target_tier, brand=brand)
    fob(db, source.material_code, 10000 + (0 if source_tier == "single" else amount), 10000,
        None if source_tier == "single" else amount)
    result = repo.initialize_sku_fobs_from_source(db, target.material_code, source.material_code)
    db.commit()
    db.expire_all()
    saved = repo.get_fob_for_country_sku(db, "CH", "NEW")
    assert result["created"] == 1
    assert saved.base_fob_eur == 10000
    assert saved.final_fob_eur == 10000 + (0 if target_tier == "single" else amount)


@pytest.mark.parametrize("reverse", [False, True])
def test_country_adjustment_uses_frozen_base(db, reverse, monkeypatch):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    single = fob(db, "S", 1000)
    dual = fob(db, "D", 1000, source="copied_from_country")
    rows = [dual, single] if reverse else [single, dual]
    monkeypatch.setattr(repo, "list_fob_by_country", lambda *_: rows)
    result = repo.adjust_country_fobs(db, "CH", 200)
    db.commit()
    assert result["adjusted"] == 2
    assert single.final_fob_eur == 1200
    assert dual.base_fob_eur == 1200
    assert dual.final_fob_eur == 1500


def test_audit_apply_reload_is_idempotent_and_matches_matrix(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    sku(db, "U", "", template="U**001")
    fob(db, "S", 19350)
    fob(db, "D", 19350, source="copied_from_country")
    fob(db, "U", 1000)
    db.commit()
    audit = repo.audit_colour_surcharge_reprice(db)
    assert audit["summary"]["missingTier"] == 1
    assert audit["summary"]["autoReprice"] == 1
    applied = routes.apply_colour_surcharge_reprice({"previewFingerprint": audit["fingerprint"]}, db, SimpleNamespace(name="test", role="admin"))
    assert applied["totals"]["updated"] == 1
    db.expire_all()
    saved = repo.get_fob_for_country_sku(db, "CH", "D")
    assert saved.final_fob_eur == 19650
    again = repo.audit_colour_surcharge_reprice(db)
    assert again["summary"]["autoReprice"] == 0
    assert repo.apply_colour_surcharge_reprice_audit(db, again["fingerprint"])["totals"]["updated"] == 0
    with pytest.raises(ValueError, match="stale"):
        repo.apply_colour_surcharge_reprice_audit(db, audit["fingerprint"])
    bom, _countries = repo.list_bom_with_fob(db)
    assert next(row for row in bom if row["materialCode"] == "D")["fobByCountry"]["CH"]["finalFobEur"] == 19650
    matrix = service.build_matrix(db, "CH", 2026)
    assert next(row for row in matrix["rows"] if row["materialCode"] == "D")["fobEur"] == 19650


def test_matrix_selection_date_uses_country_template_base_and_saved_tier(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000)
    fob(db, "D", 1300, base=1000, surcharge=300)
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 7, 15),
        valid_to=date(2026, 7, 31),
        base_fob_eur=1200,
        remark="July update",
        changed_by="test",
    )
    db.commit()

    matrix = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 7, 20))
    rows = {row["materialCode"]: row for row in matrix["rows"]}

    assert rows["S"]["fobEur"] == 1200
    assert rows["D"]["fobEur"] == 1500
    assert rows["D"]["fobPeriod"]["colourTier"] == "dual"
    assert rows["D"]["fobPeriod"]["surchargeEur"] == 300


def test_country_template_periods_remain_country_specific(db):
    sku(db, "S", "single")
    fob(db, "S", 1000, country="CH")
    fob(db, "S", 900, country="SK")
    for country, base in (("CH", 1200), ("SK", 950)):
        repo.save_country_template_fob_period(
            db,
            country_code=country,
            bom_template="T**001",
            valid_from=date(2026, 7, 1),
            valid_to=date(2026, 7, 31),
            base_fob_eur=base,
            remark=None,
            changed_by="test",
        )
    db.commit()

    ch = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 7, 20))
    sk = service.build_matrix(db, "SK", 2026, selection_date=date(2026, 7, 20))

    assert next(row for row in ch["rows"] if row["materialCode"] == "S")["fobEur"] == 1200
    assert next(row for row in sk["rows"] if row["materialCode"] == "S")["fobEur"] == 950


def test_zero_period_marks_sku_stopped_without_default_fallback(db):
    sku(db, "S", "single")
    fob(db, "S", 1000)
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 8, 1),
        valid_to=date(2026, 8, 31),
        base_fob_eur=0,
        remark="Stopped",
        changed_by="test",
    )
    db.commit()

    matrix = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 8, 10))
    row = next(item for item in matrix["rows"] if item["materialCode"] == "S")

    assert row["fobEur"] is None
    assert row["fobPeriod"]["status"] == "stopped"


def test_configured_period_gap_never_falls_back_to_legacy_fob(db):
    sku(db, "S", "single")
    fob(db, "S", 1000)
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 1, 31),
        base_fob_eur=1100,
        remark=None,
        changed_by="test",
    )
    db.commit()

    matrix = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 2, 10))
    row = next(item for item in matrix["rows"] if item["materialCode"] == "S")

    assert row["fobEur"] is None
    assert row["fobPeriod"]["status"] == "no_price"
    assert row["fobConflict"]["reason"] == "No FOB is available on 2026-02-10"


def test_no_period_schedule_keeps_undated_legacy_fob(db):
    sku(db, "S", "single")
    fob(db, "S", 1000)
    db.commit()

    matrix = service.build_matrix(db, "CH", 2026, selection_date=date(2026, 2, 10))
    row = next(item for item in matrix["rows"] if item["materialCode"] == "S")

    assert row["fobEur"] == 1000
    assert row["fobPeriod"] is None


def test_deleting_last_period_previews_default_and_retains_history(db):
    single = sku(db, "S", "single")
    fob(db, "S", 1000, base=1000)
    row = repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    db.commit()
    preview = routes.delete_bom_template_fob_period(row.country_template_fob_period_id, row.row_version,
        preview_only=True, fingerprint=None, session=db, user=SimpleNamespace(name="test", role="admin"))
    assert preview["defaultBaseFobEur"] == 1000 and preview["lastPeriod"] is True
    with pytest.raises(HTTPException) as exc:
        routes.delete_bom_template_fob_period(row.country_template_fob_period_id, row.row_version,
            preview_only=False, fingerprint="stale", session=db, user=SimpleNamespace(name="test", role="admin"))
    assert exc.value.status_code == 409
    confirmed = routes.delete_bom_template_fob_period(row.country_template_fob_period_id, row.row_version,
        preview_only=False, fingerprint=preview["fingerprint"], session=db, user=SimpleNamespace(name="test", role="admin"))
    assert confirmed["deleted"] is True
    db.commit(); db.expire_all()
    assert repo.list_country_template_fob_periods(db, "CH", "T**001") == []
    assert repo.has_country_template_fob_periods(db, "CH", "T**001") is False
    amount, conflict, evidence = service.resolve_date_effective_fob(db, "CH", single, date(2026, 8, 2))
    assert amount is None and conflict is None and evidence is None  # caller now uses its undated default
    assert service.build_matrix(db, "CH", 2026, selection_date=date(2026, 8, 2))["rows"][0]["fobEur"] == 1000
    assert db.get(models.CountryTemplateFobPeriod, row.country_template_fob_period_id).status == "deleted"
    saved = service.update_quantity_cell(db, "CH", 2026, 8, "S", 2, "test", 1)
    assert saved["fob_eur"] == 1000
    # The same start date can be used again without deleting audit history.
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=1300, remark=None, changed_by="test")
    db.commit()
    assert repo.has_country_template_fob_periods(db, "CH", "T**001") is True


def test_period_only_bom_country_search_includes_materials_and_keeps_raw_default(db):
    sku(db, "S", "single"); sku(db, "D", "dual")
    sku(db, "OTHER", "single", template="OTHER**001")
    fob(db, "OTHER", 900, country="NL")
    first = repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2027, 1, 1), valid_to=None, base_fob_eur=1300, remark=None, changed_by="test")
    db.commit()
    items, countries = repo.list_bom_with_fob(db, country_code="CH")
    assert countries == ["CH"]  # Explicit country search projects only the requested price columns.
    assert {item["materialCode"] for item in items} == {"S", "D"}
    assert items[0]["fobByCountry"] == {}
    assert [p["baseFobEur"] for p in items[0]["fobPeriodsByCountry"]["CH"]] == [1200, 1300]
    assert repo.list_active_fob_material_codes(db, "CH") == {"S", "D"}
    preview = repo.preview_country_template_fob_period_deletion(db, first.country_template_fob_period_id, first.row_version)
    assert preview["lastPeriod"] is False
    repo.delete_country_template_fob_period(db, first.country_template_fob_period_id, first.row_version)
    db.commit()
    assert service.resolve_date_effective_fob(db, "CH", repo.get_sku_by_material_code(db, "S"), date(2026, 8, 2))[0] is None


@pytest.mark.parametrize("default", [None, 0, 1000])
def test_last_period_deletion_returns_to_default_or_no_price_without_restore(db, default):
    sku(db, "S", "single")
    if default is not None:
        fob(db, "S", default, base=default)
    row = repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    db.commit()
    preview = repo.preview_country_template_fob_period_deletion(db, row.country_template_fob_period_id, row.row_version)
    assert preview["defaultBaseFobEur"] == default
    routes.delete_bom_template_fob_period(row.country_template_fob_period_id, row.row_version,
        preview_only=False, fingerprint=preview["fingerprint"], session=db, user=SimpleNamespace(name="test", role="admin"))
    assert repo.has_country_template_fob_periods(db, "CH", "T**001") is False
    if not default:
        with pytest.raises(HTTPException) as error:
            vehicle_service._resolve_pi_pricing_date(db, "CH", "S", 2026, 8, date(2026, 8, 1))
        assert error.value.status_code == 409
    else:
        assert vehicle_service._resolve_pi_pricing_date(db, "CH", "S", 2026, 8, date(2026, 8, 1)) == date(2026, 8, 1)


def test_last_period_delete_does_not_pick_ambiguous_default_and_rechecks_changes(db):
    sku(db, "S", "single")
    first = fob(db, "S", 1000, base=1000)
    row = repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    db.commit()
    preview = repo.preview_country_template_fob_period_deletion(db, row.country_template_fob_period_id, row.row_version)
    first.final_fob_eur = first.base_fob_eur = 1100
    db.commit()
    with pytest.raises(HTTPException) as error:
        routes.delete_bom_template_fob_period(row.country_template_fob_period_id, row.row_version,
            preview_only=False, fingerprint=preview["fingerprint"], session=db, user=SimpleNamespace(name="test", role="admin"))
    assert error.value.status_code == 409
    assert row.status == "active"
    sku(db, "OTHER", "single")
    fob(db, "OTHER", 1300, base=1300)
    db.commit()
    with pytest.raises(ValueError, match="(?i)conflict|ambiguous"):
        repo.preview_country_template_fob_period_deletion(db, row.country_template_fob_period_id, row.row_version)
    assert row.status == "active"


def test_country_month_availability_prices_all_colours_and_keeps_other_country(db):
    sku(db, "S", "single"); sku(db, "D", "dual")
    for country in ("CH", "SK"):
        fob(db, "S", 1000, base=1000, country=country)
        fob(db, "D", 1300, base=1000, surcharge=300, country=country)
    for month, price in ((8, 1200), (12, 1500)):
        repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
            valid_from=date(2026, month, 14 if month == 12 else 1), valid_to=None if month == 12 else date(2026, 8, 31),
            base_fob_eur=price, remark=None, changed_by="test")
    db.commit()
    ch = {row["materialCode"]: row for row in service.build_matrix(db, "CH", 2026)["rows"]}
    sk = {row["materialCode"]: row for row in service.build_matrix(db, "SK", 2026)["rows"]}
    assert ch["S"]["months"]["8"]["fobEur"] == 1200
    assert ch["D"]["months"]["8"]["fobEur"] == 1500
    assert ch["D"]["months"]["9"]["isEditable"] is False
    assert ch["D"]["months"]["12"]["availableRanges"] == [{"from": "2026-12-14", "to": "2026-12-31"}]
    assert ch["D"]["months"]["12"]["requiresOrderDate"] is True
    assert sk["D"]["months"]["9"]["fobEur"] == 1300
    with pytest.raises(ValueError, match="No country FOB"):
        service.update_quantity_cell(db, "CH", 2026, 9, "D", 2, "test", 1)
    saved = service.update_quantity_cell(db, "CH", 2026, 12, "D", 2, "test", 1)
    db.commit(); db.expire_all()
    assert saved["fob_eur"] == 1800
    cell = repo.list_quantities_for_country_year(db, "CH", 2026)[0]
    assert cell.fob_eur == 1800


def test_period_only_prices_are_visible_and_pi_uses_the_same_month_and_date(db):
    sku(db, "S", "single"); sku(db, "D", "dual")
    # No old SKU FOB is needed: the country's template period is authoritative.
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 14), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    db.commit(); db.expire_all()
    rows = {row["materialCode"]: row for row in service.build_matrix(db, "CH", 2026)["rows"]}
    assert rows["D"]["months"]["8"]["fobEur"] == 1500
    with pytest.raises(HTTPException, match="choose an available orderDate"):
        vehicle_service._resolve_pi_pricing_date(db, "CH", "D", 2026, 8, None)
    for day in (14, 31):
        resolved = vehicle_service._resolve_pi_pricing_date(db, "CH", "D", 2026, 8, date(2026, 8, day))
        assert vehicle_service._line_payload_from_material(db, "CH", "D", {}, pricing_date=resolved)["fobEur"] == 1500
    with pytest.raises(HTTPException):
        vehicle_service._line_payload_from_material(db, "CH", "D", {}, pricing_date=date(2026, 8, 13))


def test_pi_full_month_same_price_periods_and_legacy_conflict(db):
    sku(db, "S", "single")
    for start, end in ((1, 14), (15, 31)):
        repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
            valid_from=date(2026, 8, start), valid_to=date(2026, 8, end), base_fob_eur=1000, remark=None, changed_by="test")
    db.commit()
    assert vehicle_service._resolve_pi_pricing_date(db, "CH", "S", 2026, 8, None) == date(2026, 8, 1)
    second = repo.list_country_template_fob_periods(db, "CH", "T**001")[1]
    second.base_fob_eur = 1200  # An existing legacy conflict, not allowed through CRUD.
    db.commit()
    with pytest.raises(HTTPException, match="Conflicting monthly"):
        vehicle_service._resolve_pi_pricing_date(db, "CH", "S", 2026, 8, date(2026, 8, 20))


def test_lifecycle_preview_rejects_inverted_dates_without_writing(db):
    single = sku(db, "S", "single")
    db.commit()
    with pytest.raises(ValueError, match="Final order date"):
        repo.preview_bom_template_lifecycle_update(db, "S", date(2026, 9, 30), date(2026, 8, 1))
    db.expire_all()
    assert single.effective_from_date is None and single.effective_to_date is None


def test_planned_template_accepts_month_plan_before_its_start_day(db):
    future = date.today() + timedelta(days=45)
    single = sku(db, "S", "single")
    single.effective_from_date = date(future.year, future.month, 14)
    fob(db, "S", 1000, base=1000)
    db.commit()
    matrix = service.build_matrix(db, "CH", future.year)
    row = next(row for row in matrix["rows"] if row["materialCode"] == "S")
    assert row["months"][str(future.month)]["isEditable"] is True
    assert row["months"][str(future.month)]["availableRanges"][0]["from"].endswith("-14")
    assert service.update_quantity_cell(db, "CH", future.year, future.month, "S", 5, "test", 1)["fob_eur"] == 1000


def test_different_bases_in_disjoint_days_of_one_month_rejected(db):
    sku(db, "S", "single")
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 10), base_fob_eur=1000, remark=None, changed_by="test")
    db.commit()
    with pytest.raises(ValueError, match="One month"):
        repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
            valid_from=date(2026, 8, 20), valid_to=date(2026, 8, 31), base_fob_eur=1200, remark=None, changed_by="test")
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 20), valid_to=date(2026, 8, 31), base_fob_eur=1000, remark=None, changed_by="test")
    db.commit()


def test_legacy_monthly_price_conflict_blocks_only_affected_country(db):
    sku(db, "S", "single"); fob(db, "S", 1000, country="CH"); fob(db, "S", 900, country="SK")
    for day, end, base in ((1, 10, 1000), (20, 31, 1200)):
        db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template="T**001",
            valid_from=date(2026, 8, day), valid_to=date(2026, 8, end), base_fob_eur=base))
    db.commit()
    ch = service.build_matrix(db, "CH", 2026)["rows"][0]["months"]["8"]
    assert ch["isEditable"] is False and "Conflicting" in ch["reason"]
    assert service.build_matrix(db, "SK", 2026)["rows"][0]["months"]["8"]["isEditable"] is True


def test_paused_period_never_adds_surcharge_or_saves_quantity(db):
    sku(db, "D", "dual"); fob(db, "D", 1300, base=1000, surcharge=300)
    repo.save_country_template_fob_period(db, country_code="CH", bom_template="T**001",
        valid_from=date(2026, 8, 1), valid_to=date(2026, 8, 31), base_fob_eur=0, remark=None, changed_by="test")
    db.commit()
    with pytest.raises(ValueError, match="paused"):
        service.update_quantity_cell(db, "CH", 2026, 8, "D", 1, "test", 1)
    assert repo.list_quantities_for_country_year(db, "CH", 2026) == []


def test_template_lifecycle_updates_every_colour_and_is_date_derived(db):
    single = sku(db, "S", "single")
    dual = sku(db, "D", "dual")
    result = repo.update_bom_template_lifecycle(
        db,
        material_code="S",
        lifecycle_status="active",
        effective_from=date(2026, 3, 14),
        effective_to=date(2026, 9, 30),
        expected_version=single.row_version,
    )
    db.commit()

    assert result["materialCodes"] == ["D", "S"]
    assert single.effective_from_date == dual.effective_from_date == date(2026, 3, 14)
    assert single.effective_to_date == dual.effective_to_date == date(2026, 9, 30)
    assert repo.resolve_effective_lifecycle_status(single, date(2026, 3, 13)) == "not_yet_active"
    assert repo.resolve_effective_lifecycle_status(single, date(2026, 3, 14)) == "phase_out"
    assert repo.resolve_effective_lifecycle_status(single, date(2026, 9, 30)) == "phase_out"
    assert repo.resolve_effective_lifecycle_status(single, date(2026, 10, 1)) == "historical"


def test_default_matrix_scope_excludes_expired_templates_until_historical_opt_in(db):
    expired = sku(db, "OLD", "single", template="OLD**001")
    planned = sku(db, "NEW", "single", template="NEW**001")
    today = date.today()
    expired.effective_from_date = today - timedelta(days=60)
    expired.effective_to_date = today - timedelta(days=1)
    expired.is_active = False
    expired.lifecycle_status = "historical"
    planned.effective_from_date = today + timedelta(days=10)
    planned.is_active = False
    db.commit()

    rows = repo.list_active_skus(
        db,
        target_date=today + timedelta(days=10),
    )

    assert {row.material_code for row in rows} == {"NEW"}


def test_matrix_historical_opt_in_exposes_editable_backfill_row_without_reactivation(db):
    expired = sku(db, "OLD", "single", template="OLD**001")
    expired.effective_from_date = date.today() - timedelta(days=60)
    expired.effective_to_date = date.today() - timedelta(days=1)
    expired.is_active = False
    expired.lifecycle_status = "historical"
    fob(db, "OLD", 18000)
    db.commit()

    default_matrix = service.build_matrix(db, "CH", date.today().year)
    backfill_matrix = service.build_matrix(
        db,
        "CH",
        date.today().year,
        include_historical=True,
    )

    assert all(row["materialCode"] != "OLD" for row in default_matrix["rows"])
    row = next(row for row in backfill_matrix["rows"] if row["materialCode"] == "OLD")
    assert row["lifecycleStatus"] == "historical"
    assert row["historicalBackfill"] is True
    assert row["editable"] is True
    assert row["priceSource"] == "undated_default"
    assert expired.lifecycle_status == "historical"
    assert expired.is_active is False


def test_historical_quantity_requires_opt_in_and_rejects_future_month(db):
    expired = sku(db, "OLD", "single", template="OLD**001")
    expired.effective_to_date = date.today() - timedelta(days=1)
    expired.is_active = False
    expired.lifecycle_status = "historical"
    fob(db, "OLD", 18000)
    db.commit()

    with pytest.raises(ValueError, match="explicit includeHistorical"):
        service.update_quantity_cell(
            db, "CH", date.today().year, date.today().month, "OLD", 2, "tester", 1,
        )

    saved = service.update_quantity_cell(
        db,
        "CH",
        date.today().year,
        date.today().month,
        "OLD",
        2,
        "tester",
        1,
        include_historical=True,
    )
    assert saved["quantity"] == 2
    assert expired.lifecycle_status == "historical"
    assert expired.is_active is False

    future_year = date.today().year + (1 if date.today().month == 12 else 0)
    future_month = 1 if date.today().month == 12 else date.today().month + 1
    with pytest.raises(ValueError, match="future month"):
        service.update_quantity_cell(
            db,
            "CH",
            future_year,
            future_month,
            "OLD",
            1,
            "tester",
            1,
            include_historical=True,
        )


def test_historical_backfill_turns_lifecycle_rejection_into_price_preview(db):
    expired = sku(db, "OLD", "single", template="OLD**001")
    expired.effective_from_date = date(2026, 1, 1)
    expired.effective_to_date = date(2026, 6, 30)
    expired.is_active = False
    expired.lifecycle_status = "historical"
    fob(db, "OLD", 18000)
    db.commit()

    blocked_fob, blocked_conflict, _ = service.resolve_date_effective_fob(
        db, "CH", expired, date(2026, 8, 15),
    )
    preview_fob, preview_conflict, preview_evidence = service.resolve_date_effective_fob(
        db,
        "CH",
        expired,
        date(2026, 8, 15),
        allow_historical_backfill=True,
    )

    assert blocked_fob is None
    assert "expired" in blocked_conflict["reason"]
    assert preview_fob is None
    assert preview_conflict is None
    assert preview_evidence is None


def test_historical_surcharge_review_compares_saved_evidence_with_current_rule(db):
    dual = sku(db, "OLD", "dual", template="OLD**001")
    saved_fob = fob(db, "OLD", 18200, base=18000, surcharge=200)

    review = service._historical_surcharge_review(db, dual, saved_fob, None)

    assert review["status"] == "changed"
    assert review["savedSurchargeEur"] == 200
    assert review["currentSurchargeEur"] == 300
    assert review["requiresConfirmation"] is True


def test_country_period_must_stay_inside_template_lifecycle(db):
    single = sku(db, "S", "single")
    sku(db, "D", "dual")
    repo.update_bom_template_lifecycle(
        db,
        material_code="S",
        lifecycle_status="active",
        effective_from=date(2026, 3, 14),
        effective_to=date(2026, 9, 30),
        expected_version=single.row_version,
    )

    with pytest.raises(ValueError, match="starts before"):
        repo.save_country_template_fob_period(
            db,
            country_code="CH",
            bom_template="T**001",
            valid_from=date(2026, 3, 1),
            valid_to=date(2026, 6, 30),
            base_fob_eur=1000,
            remark=None,
            changed_by="test",
        )
    with pytest.raises(ValueError, match="exceeds"):
        repo.save_country_template_fob_period(
            db,
            country_code="CH",
            bom_template="T**001",
            valid_from=date(2026, 3, 14),
            valid_to=None,
            base_fob_eur=1000,
            remark=None,
            changed_by="test",
        )


def test_lifecycle_preview_lists_price_periods_that_would_be_outside(db):
    sku(db, "S", "single")
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 12, 31),
        base_fob_eur=1000,
        remark=None,
        changed_by="test",
    )
    db.commit()

    preview = repo.preview_bom_template_lifecycle_update(
        db,
        "S",
        date(2026, 1, 1),
        date(2026, 9, 30),
    )

    assert preview["canApply"] is False
    assert preview["affectedPeriods"][0]["afterTemplateEnd"] is True
    with pytest.raises(ValueError, match="exceed"):
        repo.update_bom_template_lifecycle(
            db,
            material_code="S",
            lifecycle_status="phase_out",
            effective_from=date(2026, 1, 1),
            effective_to=date(2026, 9, 30),
            expected_version=1,
        )


def test_lifecycle_route_maps_legacy_months_and_updates_template_group(db):
    single = sku(db, "S", "single")
    dual = sku(db, "D", "dual")
    result = routes.patch_sku_lifecycle(
        "S",
        {
            "lifecycleStatus": "phase_out",
            "effectiveFrom": "2026-03",
            "effectiveTo": "2026-09",
            "rowVersion": single.row_version,
        },
        session=db,
        user=SimpleNamespace(name="test", role="admin"),
    )

    assert result["effectiveFrom"] == "2026-03-01"
    assert result["effectiveTo"] == "2026-09-30"
    assert result["materialCodes"] == ["D", "S"]
    assert single.effective_from_date == dual.effective_from_date == date(2026, 3, 1)
    assert single.effective_to_date == dual.effective_to_date == date(2026, 9, 30)


def test_lifecycle_route_preserves_omitted_boundary_and_clears_explicit_null(db):
    single = sku(db, "S", "single")
    single.effective_from_date = date(2026, 3, 14)
    single.effective_to_date = date(2026, 9, 30)
    db.commit()

    preserved = routes.patch_sku_lifecycle(
        "S",
        {
            "lifecycleStatus": "phase_out",
            "effectiveFrom": "2026-04-01",
            "rowVersion": single.row_version,
        },
        session=db,
        user=SimpleNamespace(name="test", role="admin"),
    )
    assert preserved["effectiveFrom"] == "2026-04-01"
    assert preserved["effectiveTo"] == "2026-09-30"

    cleared = routes.patch_sku_lifecycle(
        "S",
        {
            "lifecycleStatus": "active",
            "effectiveTo": None,
            "rowVersion": preserved["rowVersions"]["S"],
        },
        session=db,
        user=SimpleNamespace(name="test", role="admin"),
    )
    assert cleared["effectiveFrom"] == "2026-04-01"
    assert cleared["effectiveTo"] is None


def test_lifecycle_route_returns_actionable_period_conflict(db):
    single = sku(db, "S", "single")
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 12, 31),
        base_fob_eur=1000,
        remark=None,
        changed_by="test",
    )
    db.commit()

    preview = routes.patch_sku_lifecycle(
        "S",
        {
            "lifecycleStatus": "phase_out",
            "effectiveFrom": "2026-01-01",
            "effectiveTo": "2026-09-30",
            "rowVersion": single.row_version,
            "previewOnly": True,
        },
        session=db,
        user=SimpleNamespace(name="test", role="admin"),
    )
    assert preview["canApply"] is False

    with pytest.raises(routes.HTTPException) as exc:
        routes.patch_sku_lifecycle(
            "S",
            {
                "lifecycleStatus": "phase_out",
                "effectiveFrom": "2026-01-01",
                "effectiveTo": "2026-09-30",
                "rowVersion": single.row_version,
            },
            session=db,
            user=SimpleNamespace(name="test", role="admin"),
        )

    assert exc.value.status_code == 409
    assert exc.value.detail["code"] == "template_lifecycle_fob_period_conflict"
    assert exc.value.detail["preview"]["affectedPeriods"][0]["countryCode"] == "CH"


def test_template_country_fob_periods_reject_overlap(db):
    repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 9, 1),
        valid_to=date(2026, 9, 15),
        base_fob_eur=1000,
        remark=None,
        changed_by="test",
    )
    with pytest.raises(ValueError, match="cannot overlap"):
        repo.save_country_template_fob_period(
            db,
            country_code="CH",
            bom_template="T**001",
            valid_from=date(2026, 9, 15),
            valid_to=date(2026, 9, 30),
            base_fob_eur=1200,
            remark=None,
            changed_by="test",
        )


def test_template_country_fob_period_update_keeps_original_scope(db, monkeypatch):
    period = repo.save_country_template_fob_period(
        db,
        country_code="CH",
        bom_template="T**001",
        valid_from=date(2026, 10, 1),
        valid_to=date(2026, 10, 31),
        base_fob_eur=1000,
        remark=None,
        changed_by="test",
    )
    db.commit()
    calls = []
    monkeypatch.setattr(
        routes,
        "validate_country_access",
        lambda _session, _name, _role, country: calls.append(country),
    )

    with pytest.raises(routes.HTTPException) as exc:
        routes.put_bom_template_fob_period(
            body={
                "periodId": str(period.country_template_fob_period_id),
                "rowVersion": period.row_version,
                "countryCode": "RO",
                "bomTemplate": "T**001",
                "validFrom": "2026-10-01",
                "validTo": "2026-10-31",
                "baseFobEur": 1100,
            },
            session=db,
            user=SimpleNamespace(name="tester", role="admin"),
        )

    assert exc.value.status_code == 400
    assert calls == ["RO", "CH"]
    db.refresh(period)
    assert period.country_code == "CH"
    assert float(period.base_fob_eur) == 1000


def test_conflicting_single_payment_terms_remain_ambiguous(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000, term="TT")
    fob(db, "S", 1200, term="LC90")
    dual = fob(db, "D", 1000, base=1000, source="manual_edit")
    audit = repo.audit_colour_surcharge_reprice(db)
    assert audit["summary"]["ambiguousBase"] == 1
    assert audit["duplicateCountryPrices"][0]["status"] == "conflict"
    assert repo.reprice_sku_colour_surcharge_fobs(db, "D")["skippedAmbiguous"] == 1
    assert dual.final_fob_eur == 1000
    with pytest.raises(repo.CountryFobConflict):
        repo.get_fob_for_country_sku(db, "CH", "S")
    fob_map, conflicts = repo.list_fobs_for_country_material_codes(
        db, "CH", ["S"], include_conflicts=True,
    )
    assert fob_map == {}
    assert conflicts[0]["materialCode"] == "S"


def test_no_single_requires_consistent_template_base_evidence(db):
    sku(db, "D", "dual")
    row = fob(db, "D", 1300, base=1000, surcharge=300, source="manual_edit")
    assert repo.audit_colour_surcharge_reprice(db)["summary"]["missingBase"] == 1
    row.fob_source_mode = "template_base"
    assert repo.audit_colour_surcharge_reprice(db)["summary"]["alreadyCorrect"] == 1
    sku(db, "E", "special")
    fob(db, "E", 1500, base=1200)
    assert repo.audit_colour_surcharge_reprice(db)["summary"]["ambiguousBase"] == 2


@pytest.mark.parametrize("name", ["Matte black", "Black & White"])
def test_explicit_single_name_never_changes_pricing_tier(db, name):
    matte = sku(db, "CP", "single", "OMODA", "BLACK**001", "OMODA9 SHS", name)
    fob(db, "CP", 26500, base=26500)
    repo.upsert_special_colour_surcharge(db, "OMODA", "CP", 300, model_name="OMODA9 SHS")
    repo.reprice_sku_colour_surcharge_fobs(db, matte.material_code)
    assert repo.get_fob_for_country_sku(db, "CH", "CP").final_fob_eur == 26500


def test_missing_rule_blocks_create_but_explicit_zero_is_valid(db):
    sku(db, "S", "single", brand="OTHER")
    sku(db, "D", "dual", brand="OTHER")
    fob(db, "S", 1000)
    assert repo.initialize_sku_fobs_from_source(db, "D", "S")["skippedMissingRule"] == 1
    assert repo.get_fob_for_country_sku(db, "CH", "D") is None
    db.add(models.BrandColourSurchargeRule(brand="OTHER", colour_type="dual", surcharge_eur=0))
    assert repo.initialize_sku_fobs_from_source(db, "D", "S")["created"] == 1
    assert repo.get_fob_for_country_sku(db, "CH", "D").final_fob_eur == 1000


def test_colour_tier_patch_requires_explicit_tier(db):
    with pytest.raises(routes.HTTPException) as error:
        routes.patch_sku_colour_tier("D", {}, db, SimpleNamespace(name="test", role="admin"))
    assert error.value.status_code == 400
    assert error.value.detail == "colourTier is required"


def test_publish_preserves_single_matte_despite_higher_imported_fob(db):
    db.add(models.CountryPaymentTermMaster(country_code="CH", country_name="Switzerland",
        payment_term_code="TT", payment_method="TT", lc_days=0))
    parsed = {"rows": [
        {"material_code": "CP", "bom_template": "BLACK**001", "brand": "OMODA",
         "model_name": "OMODA9 SHS", "version": "Premium", "colour_tier": "single",
         "exterior_color_type": "single", "exterior_color_name": "Matte black (Black Edition)",
         "exterior_color_code": "CP", "country_fobs": {"Switzerland": 26500}},
        {"material_code": "BW", "bom_template": "REGULAR**001", "brand": "OMODA",
         "model_name": "OMODA9 SHS", "version": "Premium", "colour_tier": "single",
         "exterior_color_type": "single", "exterior_color_name": "White",
         "exterior_color_code": "BW", "country_fobs": {"Switzerland": 26000}},
    ]}
    service.publish_baseline(db, parsed, "test-import", None, "test")
    db.commit()
    assert repo.get_sku_by_material_code(db, "CP").colour_tier == "single"
    assert repo.get_fob_for_country_sku(db, "CH", "CP").final_fob_eur == 26500


def test_rule_priority_tier_and_zero_with_real_records(db):
    car = sku(db, "UE", "special", "OMODA", model="OMODA9 SHS", name="Matte gray")
    fob(db, "UE", 10000, 10000)
    repo.upsert_special_colour_surcharge(db, "OMODA", "UE", 250)
    repo.upsert_special_colour_surcharge(db, "OMODA", "UE", 300, model_name="OMODA9 SHS")
    assert repo.resolve_colour_surcharge_for_sku(db, car, "special")["amount"] == 300
    assert repo.resolve_colour_surcharge_for_sku(db, car, "dual")["amount"] == 200
    assert repo.resolve_colour_surcharge_for_sku(db, car, "single")["amount"] == 0
    repo.upsert_special_colour_surcharge(db, "OMODA", "UE", 0, model_name="OMODA9 SHS")
    assert repo.resolve_colour_surcharge_for_sku(db, car, "special")["status"] == "explicit_zero"


def test_apply_rejects_changed_rule_and_route_rolls_back(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000)
    fob(db, "D", 1000)
    db.commit()
    audit = repo.audit_colour_surcharge_reprice(db)
    rule = repo.get_brand_colour_surcharge(db, "JAECOO", "dual")
    rule.surcharge_eur = 400
    with pytest.raises(routes.HTTPException) as error:
        routes.apply_colour_surcharge_reprice({"previewFingerprint": audit["fingerprint"]}, db, SimpleNamespace(name="test", role="admin"))
    assert error.value.status_code == 409
    assert repo.get_fob_for_country_sku(db, "CH", "D").final_fob_eur == 1000
    assert repo.get_brand_colour_surcharge(db, "JAECOO", "dual").surcharge_eur == 300


def test_audit_exposes_charged_single_and_existing_tier_correction_preserves_prices(db):
    sku(db, "BW", "single", name="Matte gray")
    sku(db, "UE", "single", name="Colour")
    sku(db, "ZE", "dual")
    for country, base in [("CZ", 27850), ("AT", 28000)]:
        fob(db, "BW", base, base=base, country=country)
        fob(db, "UE", base + 300, base=base, surcharge=300, country=country)
        fob(db, "ZE", base + 300, base=base, surcharge=300, country=country)
    db.commit()
    before = {(row.material_code, row.country_code): float(row.final_fob_eur)
              for row in db.scalars(select(models.CountrySkuFobResolved))}
    audit = repo.audit_colour_surcharge_reprice(db)
    assert audit["summary"]["singleSurchargeConflict"] == 2
    assert audit["summary"]["ambiguousBase"] == 2
    charged = [item for item in audit["items"] if item["materialCode"] == "UE"]
    assert all(item["category"] == "single_surcharge_conflict" for item in charged)
    assert all(item["expectedFinalFobEur"] is None for item in charged)
    assert all(item["category"] == "not_applicable"
               for item in audit["items"] if item["materialCode"] == "BW")
    apply = repo.apply_colour_surcharge_reprice_audit(db, audit["fingerprint"])
    assert apply["totals"]["updated"] == 0
    result = routes.patch_sku_colour_tier(
        "UE", {"colourTier": "special"}, db, SimpleNamespace(name="test", role="admin"),
    )
    assert result["colourTier"] == "special"
    after = {(row.material_code, row.country_code): float(row.final_fob_eur)
             for row in db.scalars(select(models.CountrySkuFobResolved))}
    assert after == before
    corrected = repo.audit_colour_surcharge_reprice(db)
    assert corrected["summary"]["singleSurchargeConflict"] == 0
    assert corrected["summary"]["ambiguousBase"] == 0
    assert corrected["summary"]["autoReprice"] == 0
    scoped = repo.audit_colour_surcharge_reprice(
        db, allowed_brands={"JAECOO"}, allowed_countries={"CZ"},
    )
    assert scoped["summary"]["rows"] == 3
    assert {item["countryCode"] for item in scoped["items"]} == {"CZ"}


def test_audit_single_surcharge_conflict_respects_authorized_scope(db):
    sku(db, "J", "single")
    sku(db, "O", "single", brand="OMODA")
    fob(db, "J", 1300, 1000, 300, country="CZ")
    fob(db, "J", 1400, 1100, 300, country="CH")
    fob(db, "O", 1200, 1000, 200, country="CZ")
    audit = repo.audit_colour_surcharge_reprice(
        db, allowed_brands={"JAECOO"}, allowed_countries={"CZ"},
    )
    assert audit["summary"]["singleSurchargeConflict"] == 1
    assert [(item["materialCode"], item["countryCode"]) for item in audit["items"]] == [("J", "CZ")]


@pytest.mark.parametrize("overwrite", [False, True])
def test_country_copy_uses_target_base_or_replaces_whole_base(db, overwrite):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000, country="CZ")
    fob(db, "D", 1000, country="CZ", source="uploaded_final_fob")
    fob(db, "S", 2000, country="CH")
    result = repo.copy_country_fobs(db, "CZ", "CH", overwrite_existing=overwrite)
    db.commit()
    expected_base = 1000 if overwrite else 2000
    assert repo.get_fob_for_country_sku(db, "CH", "D").final_fob_eur == expected_base + 300
    assert repo.get_fob_for_country_sku(db, "CH", "S").final_fob_eur == expected_base
    assert result["repriced"] == 1


def test_country_adjust_skips_conflicting_single_and_derived_rows(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000)
    fob(db, "S", 1200, term="LC90")
    fob(db, "D", 1300, 1000, 300, source="manual_edit")
    result = repo.adjust_country_fobs(db, "CH", 200)
    assert result["skippedAmbiguous"] == 3
    assert result["adjusted"] == 0
    assert db.scalars(select(models.FobResolvedHistory)).all() == []


def test_equal_payment_terms_do_not_choose_or_adjust_different_prices(db):
    sku(db, "S", "single")
    sku(db, "D", "dual")
    fob(db, "S", 1000, term="TT")
    fob(db, "S", 1000, term="LC90")
    fob(db, "D", 1000, term="TT")
    fob(db, "D", 1000, term="LC90")
    db.commit()
    audit = repo.audit_colour_surcharge_reprice(db)
    result = repo.apply_colour_surcharge_reprice_audit(db, audit["fingerprint"])
    assert result["totals"]["requested"] == 1
    assert result["totals"]["updated"] == 2
    assert repo.get_fob_for_country_sku(db, "CH", "D", "TT").final_fob_eur == 1300
    assert repo.get_fob_for_country_sku(db, "CH", "D", "LC90").final_fob_eur == 1300


def test_upsert_bom_price_normalizes_all_payment_term_rows(db):
    sku(db, "D", "dual")
    fob(db, "D", 1300, base=1000, surcharge=300, term="TT")
    fob(db, "D", 1500, base=1200, surcharge=300, term="LC90")
    replacement = models.CountrySkuFobResolved(
        baseline_version_id=db.info["baseline"],
        material_code="D", country_code="CH", payment_term_code="TT",
        base_fob_eur=1100, uploaded_fob_eur=1100,
        colour_surcharge_eur=300, final_fob_eur=1400,
        fob_source_mode="template_base",
    )
    repo.upsert_fob_resolved(db, replacement)
    db.flush()
    rows = db.scalars(select(models.CountrySkuFobResolved).where(
        models.CountrySkuFobResolved.material_code == "D",
        models.CountrySkuFobResolved.country_code == "CH",
    )).all()
    assert {(row.payment_term_code, float(row.base_fob_eur), float(row.final_fob_eur)) for row in rows} == {
        ("TT", 1100, 1400), ("LC90", 1100, 1400),
    }


def test_bom_and_matrix_keep_normal_rows_when_one_country_group_conflicts(db):
    sku(db, "BAD", "single", template="BAD**001")
    sku(db, "OK", "single", template="OK**001")
    fob(db, "BAD", 1000, term="TT")
    fob(db, "BAD", 1200, term="LC90")
    fob(db, "OK", 900)
    items, _countries, conflicts = repo.list_bom_with_fob(db, include_conflicts=True)
    bad = next(item for item in items if item["materialCode"] == "BAD")
    good = next(item for item in items if item["materialCode"] == "OK")
    assert bad["fobByCountry"]["CH"]["status"] == "conflict"
    assert good["fobByCountry"]["CH"]["finalFobEur"] == 900
    assert conflicts[0]["materialCode"] == "BAD"

    matrix = service.build_matrix(db, "CH", 2026)
    matrix_rows = {row["materialCode"]: row for row in matrix["rows"]}
    assert matrix_rows["BAD"]["fobEur"] is None
    assert matrix_rows["BAD"]["fobConflict"]["status"] == "conflict"
    assert matrix_rows["OK"]["fobEur"] == 900
    with pytest.raises(ValueError, match="Export blocked"):
        service.export_matrix(db, "CH", 2026)


def test_copy_country_conflict_isolated_to_material_country_group(db):
    sku(db, "BAD", "single", template="BAD**001")
    sku(db, "OK", "single", template="OK**001")
    fob(db, "BAD", 1000, country="CZ", term="TT")
    fob(db, "BAD", 1200, country="CZ", term="LC90")
    fob(db, "OK", 1000, country="CZ")

    result = repo.copy_country_fobs(db, "CZ", "CH")
    db.flush()

    assert result["skippedAmbiguous"] == 1
    assert repo.get_fob_for_country_sku(db, "CH", "BAD") is None
    assert repo.get_fob_for_country_sku(db, "CH", "OK").final_fob_eur == 1000


def test_copy_country_rederives_conflicting_colour_outputs_from_single_base(db):
    sku(db, "S", "single", template="T**001")
    sku(db, "D", "dual", template="T**001")
    fob(db, "S", 1000, country="CZ")
    fob(db, "D", 1200, base=1000, surcharge=200, country="CZ", term="TT")
    fob(db, "D", 1400, base=1100, surcharge=300, country="CZ", term="LC90")

    result = repo.copy_country_fobs(db, "CZ", "CH")
    db.flush()

    assert result["skippedAmbiguous"] == 0
    assert repo.get_fob_for_country_sku(db, "CH", "S").final_fob_eur == 1000
    copied_dual = repo.get_fob_for_country_sku(db, "CH", "D")
    assert copied_dual.base_fob_eur == 1000
    assert copied_dual.colour_surcharge_eur == 300
    assert copied_dual.final_fob_eur == 1300


def test_publish_requires_explicit_colour_tier(db):
    parsed = {"rows": [{
        "material_code": "NO-TIER",
        "bom_template": "NO-TIER**001",
        "brand": "OMODA",
        "model_name": "OMODA5",
        "version": "Premium",
        "exterior_color_type": "",
        "colour_tier": "",
        "exterior_color_name": "Unknown",
        "exterior_color_code": "XX",
        "country_fobs": {},
    }]}
    with pytest.raises(ValueError, match="Colour tier is required"):
        service.publish_baseline(db, parsed, "missing-tier", None, "test")


def test_country_fob_fallback_is_country_only_and_rejects_conflicting_sources(db):
    db.add_all([
        models.CountryFobSourceMapping(
            target_country_code="RO",
            target_payment_term_code="TT",
            source_country_code="HR",
            is_active=True,
        ),
        models.CountryFobSourceMapping(
            target_country_code="RO",
            target_payment_term_code="LC90",
            source_country_code="HR",
            is_active=True,
        ),
    ])
    db.flush()

    assert repo.get_country_fob_source_mapping(db, "ro", "LC90") == "HR"

    db.add(models.CountryFobSourceMapping(
        target_country_code="RO",
        target_payment_term_code="LC180",
        source_country_code="CZ",
        is_active=True,
    ))
    db.flush()
    with pytest.raises(ValueError, match="Conflicting FOB source mappings"):
        repo.get_country_fob_source_mapping(db, "RO", "TT")


def timeline_link(db, source, target, country="CH"):
    db.add(models.MaterialLifecycle(country_code=country, material_code=source,
        product_identity="JAECOO|Model|Version|BEV", valid_from=date.today() - timedelta(days=90),
        replaced_by_code=target, lifecycle_status="replaced", remark="Certificate period is not an order date"))
    db.flush()


def test_timeline_country_brand_scope_and_no_model_guess(db):
    sku(db, "A", "single", template="OLD")
    sku(db, "B", "single", template="NEXT")
    sku(db, "OMODA", "single", brand="OMODA", template="OTHER BRAND")
    fob(db, "A", 1000, base=1000, country="CH")
    fob(db, "B", 1000, base=1000, country="SE")
    fob(db, "OMODA", 1000, base=1000, country="CH")
    result = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"})
    assert [node["bomTemplate"] for node in result["nodes"]] == ["OLD"]
    assert result["links"] == result["events"] == []
    assert service.build_material_timeline(db, "CH", allowed_brands=set())["nodes"] == []
    assert len(service.build_material_timeline(db, "CH", allowed_brands=None)["nodes"]) == 2


def test_timeline_explicit_chain_includes_history_and_unquoted_successor_only_in_its_country(db):
    old = sku(db, "A", "single", template="OLD")
    old.lifecycle_status = "historical"
    old.is_active = False
    sku(db, "B", "single", template="MID")
    sku(db, "C", "single", template="NEXT")
    fob(db, "A", 1000, base=1000, country="CH")
    fob(db, "A", 1000, base=1000, country="SE")
    timeline_link(db, "A", "B")
    timeline_link(db, "B", "C")
    result = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"})
    assert {node["bomTemplate"] for node in result["nodes"]} == {"OLD", "MID", "NEXT"}
    assert len(result["links"]) == 2
    assert all(node["effectiveFrom"] is None and node["effectiveTo"] is None for node in result["nodes"])
    assert result["events"] == []  # Never borrow the lifecycle/certificate evidence date.
    assert all(node["quotedSkuCount"] == 0 for node in result["nodes"])
    other = service.build_material_timeline(db, "SE", allowed_brands={"JAECOO"})
    assert [node["bomTemplate"] for node in other["nodes"]] == ["OLD"]


def test_timeline_dates_events_are_per_base_and_read_only(db):
    from sqlalchemy import event
    first = sku(db, "A", "single")
    second = sku(db, "B", "dual")
    for value in (first, second):
        value.effective_from_date = date.today() - timedelta(days=2)
        value.effective_to_date = date.today() + timedelta(days=3)
        value.remark = "Upgrade note"
    fob(db, "A", 1000, base=1000)
    fob(db, "B", 1300, base=1000, surcharge=300)
    db.commit()
    statements = []
    def capture(_conn, _cursor, statement, _params, _context, _many):
        statements.append(statement.lstrip().split()[0].upper())
    engine = db.get_bind()
    event.listen(engine, "before_cursor_execute", capture)
    try:
        result = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"})
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    assert len(result["nodes"]) == 1 and len(result["events"]) == 2
    assert result["nodes"][0]["quotedSkuCount"] == 2
    assert result["nodes"][0]["notes"] == ["Upgrade note"]
    assert not set(statements) & {"INSERT", "UPDATE", "DELETE"}
    assert not db.new and not db.dirty


@pytest.mark.parametrize("offset,included", [(-31, False), (-30, True), (30, True), (31, False)])
def test_timeline_event_window_boundaries(db, offset, included):
    value = sku(db, "A", "single")
    value.effective_from_date = date.today() + timedelta(days=offset)
    fob(db, "A", 1000, base=1000)
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert bool(result["events"]) == included


def test_timeline_country_period_gap_does_not_fall_back_to_default(db):
    value = sku(db, "A", "single")
    fob(db, "A", 1000, base=1000)
    today = date.today()
    db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template=value.bom_template,
        valid_from=today + timedelta(days=5), base_fob_eur=2000, status="active"))
    db.flush()
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert result["nodes"][0]["quotedSkuCount"] == 0
    assert result["nodes"][0]["quoteDate"] == today.isoformat()
    assert result["nodes"][0]["quoteReasons"]


def test_timeline_metadata_conflict_stops_events_and_quote_claim(db):
    first = sku(db, "A", "single")
    second = sku(db, "B", "single")
    first.effective_from_date = date.today()
    second.effective_from_date = date.today() + timedelta(days=1)
    fob(db, "A", 1000, base=1000)
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    node = result["nodes"][0]
    assert node["effectiveFrom"] is None and node["issues"]
    assert node["quotedSkuCount"] == 0 and result["events"] == []


@pytest.mark.parametrize("problem", ["cycle", "partial", "other_brand"])
def test_timeline_does_not_turn_invalid_sku_relations_into_base_chains(db, problem):
    sku(db, "A", "single", template="OLD")
    sku(db, "B", "single", template="NEXT", brand="OMODA" if problem == "other_brand" else "JAECOO")
    fob(db, "A", 1000, base=1000)
    timeline_link(db, "A", "B")
    if problem == "cycle":
        timeline_link(db, "B", "A")
    if problem == "partial":
        sku(db, "A2", "single", template="OLD")
        fob(db, "A2", 1000, base=1000)
    result = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"})
    assert result["links"] == []
    assert [node["bomTemplate"] for node in result["nodes"]] == ["OLD"]
    assert result["nodes"][0]["issues"]
    assert "OMODA" not in str(result)


@pytest.mark.parametrize("role,country,brands,expected", [
    ("viewer", "SE", ["JAECOO"], 403),
    ("order_filler", "CH", ["JAECOO"], 403),
    ("order_filler", "SE", ["JAECOO"], 200),
    ("order_filler", "SE", [], 200),
    ("editor", "SE", ["JAECOO"], 200),
    ("admin", "CH", [], 200),
    ("developer", "CH", [], 200),
])
def test_timeline_http_uses_saved_country_brand_and_existing_role_guard(monkeypatch, role, country, brands, expected):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.core.security import UserContext, get_current_user
    from app.db.session import get_db_session
    account = SimpleNamespace(primary_country_code="SE", secondary_country_codes=[], brands=brands)
    class AccountQuery:
        def filter(self, *_args):
            return self
        def first(self):
            return account
    session = SimpleNamespace(query=lambda _model: AccountQuery())
    calls = []
    def read(_session, country_code, *, allowed_brands):
        calls.append((country_code, allowed_brands))
        return {"countryCode": country_code, "nodes": []}
    monkeypatch.setattr(routes, "build_material_timeline", read)
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[get_current_user] = lambda: UserContext(name="test", role=role)
    app.dependency_overrides[get_db_session] = lambda: session
    with TestClient(app) as client:
        response = client.get(f"/order-genius/material-timeline?country={country}&brand=OMODA")
    assert response.status_code == expected
    if expected == 200:
        assert calls == [(country, None if role in {"admin", "developer"} else set(brands))]
    else:
        assert calls == []


@pytest.mark.parametrize("role,expected", [("viewer",403),("order_filler",403),("editor",403),("admin",200),("developer",200)])
def test_wvta_reference_upload_uses_existing_admin_role_guard(monkeypatch, role, expected):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.core.security import UserContext, get_current_user
    calls = []
    def save(upload_id, *, triggered_by, expected_sha256):
        calls.append((upload_id, triggered_by, expected_sha256))
        return {"recordCount":1}
    monkeypatch.setattr(routes, "set_wvta_reference", save)
    app = FastAPI(); app.include_router(routes.router)
    app.dependency_overrides[get_current_user] = lambda: UserContext(name="test", role=role)
    with TestClient(app) as client:
        response = client.post("/order-genius/wvta-reference", json={"uploadId":"source", "sha256":"digest"})
    assert response.status_code == expected
    assert calls == ([("source","test","digest")] if expected == 200 else [])


def test_timeline_wvta_opt_in_keeps_bom_projection_unchanged(db, monkeypatch):
    value = sku(db, "A", "single", template="T7000Z5**MY0031")
    value.effective_from = date(2026,10,20); value.effective_to = date(2026,11,1)
    fob(db, "A", 1000, base=1000); db.commit()
    calls = []
    def read(country, nodes):
        calls.append((country, [node["bomTemplate"] for node in nodes]))
        return {"status":"ready", "records":[{"validFrom":"2026-03-11"}]}
    monkeypatch.setattr(routes, "read_wvta_reference", read)
    user = SimpleNamespace(name="admin", role="admin")
    plain = routes.get_material_timeline("CH", db, user)
    assert calls == []
    overlay = routes.get_material_timeline("CH", db, user, True)
    assert calls == [("CH", [value.bom_template])]
    assert {key:val for key,val in overlay.items() if key != "wvtaReference"} == plain
    assert value.effective_from == date(2026,10,20) and not db.dirty


def test_timeline_deleted_period_resumes_existing_default_and_future_quote_uses_first_order_day(db):
    value = sku(db, "A", "single")
    value.effective_from_date = date.today() + timedelta(days=15)
    fob(db, "A", 1000, base=1000)
    db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template=value.bom_template,
        valid_from=date.today(), base_fob_eur=0, status="deleted"))
    db.flush()
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert result["nodes"][0]["quotedSkuCount"] == 1
    assert result["nodes"][0]["quoteDate"] == value.effective_from_date.isoformat()
    assert result["nodes"][0]["lifecycleStatus"] == "not_yet_active"


def test_timeline_zero_period_never_becomes_a_positive_quote(db):
    value = sku(db, "A", "dual")
    fob(db, "A", 1300, base=1000, surcharge=300)
    db.add(models.CountryTemplateFobPeriod(country_code="CH", bom_template=value.bom_template,
        valid_from=date.today(), base_fob_eur=0, status="active"))
    db.flush()
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert result["nodes"][0]["quotedSkuCount"] == 0


def test_timeline_base_detail_does_not_reveal_other_markets_colour_skus(db):
    sku(db, "CH-COLOUR", "single")
    sku(db, "NL-ONLY-COLOUR", "single")
    fob(db, "CH-COLOUR", 1000, base=1000)
    fob(db, "NL-ONLY-COLOUR", 1000, base=1000, country="NL")
    result = service.build_material_timeline(db, "CH", allowed_brands={"JAECOO"})
    assert result["nodes"][0]["materialCodes"] == ["CH-COLOUR"]
    assert "NL-ONLY-COLOUR" not in str(result)


def test_timeline_incomplete_relation_does_not_suppress_trustworthy_dates(db):
    value = sku(db, "A", "single")
    value.effective_to_date = date.today() + timedelta(days=1)
    fob(db, "A", 1000, base=1000)
    timeline_link(db, "A", "MISSING")
    result = service.build_material_timeline(db, "CH", allowed_brands=None)
    assert result["nodes"][0]["issues"]
    assert result["events"] == [{"bomTemplate": "T**001", "kind": "end", "date": value.effective_to_date.isoformat()}]
