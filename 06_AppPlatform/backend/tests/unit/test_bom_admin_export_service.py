from __future__ import annotations

import openpyxl

from app.services.bom_admin_export_service import (
    BOM_ADMIN_EXPORT_SCHEMA_VERSION,
    build_bom_admin_workbook,
)


def _item(
    material_code: str,
    *,
    model: str = "JAECOO8 SHS",
    colour: str,
    code: str,
    tier: str,
    base: float,
    surcharge: float,
) -> dict:
    return {
        "materialCode": material_code,
        "modelCode": "T26SHS",
        "modelName": model,
        "version": "Luxury-AWD(5座）",
        "bomTemplate": "T6481QN**LX0002",
        "colour": colour,
        "colourCode": code,
        "colourTier": tier,
        "interiorColorName": "Black-Black",
        "fobByCountry": {
            "AT": {
                "baseFobEur": base,
                "colourSurchargeEur": surcharge,
                "finalFobEur": base + surcharge,
            },
            "CZ": {
                "baseFobEur": base - 150,
                "colourSurchargeEur": surcharge,
                "finalFobEur": base - 150 + surcharge,
            },
        },
    }


def test_bom_admin_export_matches_material_master_shape_and_preserves_tier() -> None:
    workbook_bytes = build_bom_admin_workbook(
        [
            _item("T6481QNBWLX0002", colour="Khaki white", code="BW", tier="single", base=28000, surcharge=0),
            _item("T6481QNZELX0002", colour="Black & White", code="ZE", tier="dual", base=28000, surcharge=300),
            _item("T6481QNUELX0002", colour="Matte Gray", code="UE", tier="special", base=28000, surcharge=300),
        ],
        ["AT", "CZ"],
    )

    wb = openpyxl.load_workbook(workbook_bytes, data_only=False)
    assert wb.sheetnames == ["JAECOO8 SHS", "_BOM_ADMIN_SCHEMA"]
    ws = wb["JAECOO8 SHS"]
    assert ws["A1"].value == "No."
    assert ws["E1"].value == "BOM Template"
    assert ws["F1"].value == "Material Code"
    assert ws["J1"].value == "Colour Tier"
    assert ws["K1"].value == "Austria (AT) FOB"
    assert ws["K2"].value == "Single Base"
    assert ws["L2"].value == "Surcharge"
    assert ws["M2"].value == "Final FOB"

    exported = {
        ws.cell(row, 6).value: {
            "tier": ws.cell(row, 10).value,
            "base": ws.cell(row, 11).value,
            "surcharge": ws.cell(row, 12).value,
            "final": ws.cell(row, 13).value,
            "fill": ws.cell(row, 7).fill.fgColor.rgb,
        }
        for row in range(3, 6)
    }
    assert exported["T6481QNBWLX0002"]["tier"] == "Single"
    assert exported["T6481QNZELX0002"]["tier"] == "Dual"
    assert exported["T6481QNUELX0002"]["tier"] == "Special"
    assert exported["T6481QNZELX0002"]["base"] == 28000
    assert exported["T6481QNZELX0002"]["surcharge"] == 300
    assert exported["T6481QNZELX0002"]["final"] == 28300
    assert exported["T6481QNZELX0002"]["fill"] == exported["T6481QNUELX0002"]["fill"]
    assert exported["T6481QNBWLX0002"]["fill"] != exported["T6481QNZELX0002"]["fill"]

    schema = wb["_BOM_ADMIN_SCHEMA"]
    assert schema.sheet_state == "hidden"
    assert schema["B1"].value == BOM_ADMIN_EXPORT_SCHEMA_VERSION


def test_bom_admin_export_can_limit_prices_to_one_country() -> None:
    workbook_bytes = build_bom_admin_workbook(
        [_item("T6481QNZELX0002", colour="Black & White", code="ZE", tier="dual", base=28000, surcharge=300)],
        ["AT", "CZ"],
        country_code="CZ",
    )

    wb = openpyxl.load_workbook(workbook_bytes, data_only=False)
    ws = wb["JAECOO8 SHS"]
    assert ws.max_column == 13
    assert ws["K1"].value == "Czech Republic (CZ) FOB"
    assert ws["K3"].value == 27850
    assert ws["L3"].value == 300
    assert ws["M3"].value == 28150


def test_bom_admin_export_rejects_unknown_country_scope() -> None:
    try:
        build_bom_admin_workbook(
            [_item("T6481QNZELX0002", colour="Black & White", code="ZE", tier="dual", base=28000, surcharge=300)],
            ["AT", "CZ"],
            country_code="SE",
        )
    except ValueError as exc:
        assert str(exc) == "Country SE has no active BOM FOB data"
    else:
        raise AssertionError("Expected an unavailable country to be rejected")
