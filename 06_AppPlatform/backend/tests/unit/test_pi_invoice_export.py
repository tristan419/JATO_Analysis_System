"""Synthetic fixtures only: actual customer templates stay outside this public repo."""
import json
from copy import deepcopy
from io import BytesIO
from unittest.mock import Mock

import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from app.core.security import get_current_user, UserContext

from app.api.routes import order_genius_vehicle_allocation as route
from app.services import order_genius_vehicle_exporter as exporter


@pytest.fixture
def templates(tmp_path, monkeypatch):
    profiles = {}
    for country in ("CH", "SE"):
        profiles[country] = dict(name=f"{country} OJ", buyerName="Example Buyer",
            buyerAddress="Example address in Norway", buyerEmail="", portOfShipment="Loading",
            portOfDischarge="Discharge")
        wb = openpyxl.Workbook()
        ws = wb.active
        for coordinate, value in {"A1": "Example Seller", "A3": "Seller address", "C24": "Example bank",
                                  "C21": "CIF" if country == "CH" else "CFR", "C22": "TT" if country == "CH" else "LC 90"}.items():
            ws[coordinate] = value
        ws["B28"] = "Old sample vehicle"
        ws.merge_cells("B28:B30")
        ws["A35"] = "Old sample footer"
        ws.auto_filter.ref = "A27:XFD48"
        wb.save(tmp_path / f"{country}_OJ.xlsx")
    (tmp_path / "profiles.json").write_text(json.dumps(profiles))
    monkeypatch.setattr(exporter, "PI_TEMPLATE_ROOT", tmp_path)
    return tmp_path


def pi(country="SE", quantity=2):
    return {"header": {"piCode": f"PI-{country}-202609-001", "countryCode": country,
        "orderingAccountCode": country, "orderDate": "2026-09-01"},
        "lines": [{"piLineCode": "L01", "brand": "OMODA", "quantity": quantity}],
        "vehicles": [{"piLineCode": "L01", "countryCode": country, "brand": "OMODA",
            "modelName": "MODEL BEV", "version": "Trim", "powertrain": "BEV",
            "exteriorColorName": "White", "exteriorColorCode": "BW", "interiorColorName": "Black",
            "materialCode": "SAVED-MATERIAL", "fobEur": 10000.11, "freightEur": 123.45,
            "insuranceEur": 6.78} for _ in range(quantity)]}


@pytest.mark.parametrize("country,qty_col,fob_col,total_col", [("CH", "F", "H", "L"), ("SE", "H", "J", "M")])
def test_invoice_saved_snapshots_formulas_cached_totals_and_print(templates, country, qty_col, fob_col, total_col):
    original = (templates / f"{country}_OJ.xlsx").read_bytes()
    detail = pi(country, quantity=45)
    before = deepcopy(detail)
    output = exporter.generate_pi_invoice_excel(detail, {"handlingEur": 10 if country == "CH" else 0})
    formula = openpyxl.load_workbook(output).active
    values = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    assert formula.title == "Proforma Invoice"
    assert formula["A7"].value == "PROFORMA INVOICE"
    assert "Example Buyer" in formula["B9"].value
    assert values[f"{qty_col}28"].value == 45
    assert values[f"{qty_col}29"].value == 45
    assert values[f"{fob_col}29"].value == 450004.95
    assert values[f"{total_col}29"].value == (456315.3 if country == "CH" else 455865.3)
    assert formula[f"{total_col}28"].value.startswith("=ROUND(")
    assert formula[f"{total_col}29"].value == f"=SUM({total_col}28:{total_col}28)"
    assert formula["A30"].value == "10-Port of Shipment: Loading"
    assert not any(cell.value == "Old sample footer" for row in formula for cell in row)
    assert formula.auto_filter.ref is None
    assert formula.page_setup.fitToWidth == 1
    assert formula.print_title_rows == "$26:$27"
    assert not any(merged.min_row == 28 for merged in formula.merged_cells.ranges)
    assert (templates / f"{country}_OJ.xlsx").read_bytes() == original
    assert detail == before


def test_different_saved_costs_split_rows_and_zero_override_is_explicit(templates):
    detail = pi()
    detail["vehicles"][1]["freightEur"] = 200
    split = openpyxl.load_workbook(exporter.generate_pi_invoice_excel(detail, {}), data_only=True).active
    assert (split["H28"].value, split["H29"].value, split["H30"].value) == (1, 1, 2)
    merged = openpyxl.load_workbook(exporter.generate_pi_invoice_excel(detail, {"freightEur": 0, "insuranceEur": 0}), data_only=True).active
    assert merged["H28"].value == 2
    assert merged["M29"].value == 20000.22


@pytest.mark.parametrize("country,version_col,qty_col,amount_col,total_col", [
    ("CH", "C", "F", "H", "L"), ("SE", "D", "H", "J", "M"),
])
def test_model_version_merge_only_consecutive_runs_without_changing_details(
    templates, country, version_col, qty_col, amount_col, total_col,
):
    detail = pi(country, quantity=8)
    labels = [("A", "V1"), ("A", "V1"), ("A", "V2"), ("A", "V2"),
              ("B", "V2"), ("B", "V2"), ("A", "V1"), ("A", "V1")]
    for index, (vehicle, (model, version)) in enumerate(zip(detail["vehicles"], labels)):
        vehicle.update(modelName=model, version=version, materialCode=f"MATERIAL-{index}")
    before = deepcopy(detail)
    output = exporter.generate_pi_invoice_excel(detail, {})
    ws = openpyxl.load_workbook(output).active
    values = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    body_merges = {str(merged) for merged in ws.merged_cells.ranges if 28 <= merged.min_row < 36}
    assert body_merges == {"B28:B31", "B32:B33", "B34:B35",
                          *(f"{version_col}{start}:{version_col}{start + 1}" for start in (28, 30, 32, 34))}
    for row in range(28, 36):
        assert ws[f"A{row}"].value == row - 27
        assert ws[f"{'E' if country == 'CH' else 'G'}{row}"].value == f"MATERIAL-{row - 28}"
        assert values[f"{qty_col}{row}"].value == 1
        assert values[f"{amount_col}{row}"].value == 10000.11
        if country == "SE":
            assert ws[f"C{row}"].value == "BEV"
    for merged in body_merges:
        assert ws[merged.split(":")[0]].alignment.horizontal == "center"
        assert ws[merged.split(":")[0]].alignment.vertical == "center"
    assert values[f"{qty_col}36"].value == 8
    assert values[f"{total_col}36"].value == 81042.72
    assert ws[f"{total_col}36"].value == f"=SUM({total_col}28:{total_col}35)"
    assert detail == before


def test_empty_model_and_version_are_not_merged(templates):
    detail = pi()
    for index, vehicle in enumerate(detail["vehicles"]):
        vehicle.update(modelName="", version=None, materialCode=f"MATERIAL-{index}")
    ws = openpyxl.load_workbook(exporter.generate_pi_invoice_excel(detail, {})).active
    assert not any(merged.min_row == 28 for merged in ws.merged_cells.ranges)


@pytest.mark.parametrize("key,value", [("freightEur", None), ("insuranceEur", None), ("fobEur", 0),
                                      ("fobEur", None), ("freightEur", -1), ("insuranceEur", float("nan"))])
def test_invalid_or_missing_costs_block_instead_of_guessing(templates, key, value):
    detail = pi()
    detail["vehicles"][0][key] = value
    with pytest.raises(ValueError):
        exporter.generate_pi_invoice_excel(detail, {})


def test_missing_costs_can_be_filled_for_file_without_order_mutation(templates):
    detail = pi()
    for vehicle in detail["vehicles"]:
        vehicle["freightEur"] = vehicle["insuranceEur"] = None
    context = exporter.pi_invoice_context(detail)
    assert context["missingFreightUnits"] == 2
    result = exporter.generate_pi_invoice_excel(detail, {"freightEur": 0, "insuranceEur": 0})
    assert openpyxl.load_workbook(result, data_only=True).active["M29"].value == 20000.22
    assert detail["vehicles"][0]["freightEur"] is None


def test_account_country_not_buyer_country_selects_template(templates):
    assert exporter.pi_invoice_context(pi())["templateKey"] == "SE_OJ"
    assert "Norway" in exporter.pi_invoice_context(pi())["buyerAddress"]
    detail = pi("CH")
    detail["header"]["orderingAccountCode"] = "SE"
    assert exporter.pi_invoice_context(detail)["templateKey"] == "SE_OJ"


@pytest.mark.parametrize("brand", ["CHERY", "", None])
def test_unknown_or_other_brand_not_silently_using_oj_template(templates, brand):
    detail = pi()
    detail["lines"][0]["brand"] = brand
    with pytest.raises(ValueError, match="No matching"):
        exporter.pi_invoice_context(detail)


def test_dynamic_rows_and_text_are_safe(templates):
    detail = pi(quantity=20)
    for index, vehicle in enumerate(detail["vehicles"]):
        vehicle["materialCode"] = f"MATERIAL-{index}"
    detail["vehicles"][0]["modelName"] = "=BADFORMULA()"
    ws = openpyxl.load_workbook(exporter.generate_pi_invoice_excel(detail, {"referenceNo": "=BAD()"})).active
    assert ws["B28"].data_type == "s"
    assert ws["K9"].data_type == "s"
    assert ws["M48"].value == "=SUM(M28:M47)"
    assert ws["A49"].value == "10-Port of Shipment: Loading"


def test_incomplete_saved_pi_is_not_exported_as_complete(templates):
    detail = pi()
    detail["vehicles"].pop()
    with pytest.raises(ValueError, match="incomplete"):
        exporter.generate_pi_invoice_excel(detail, {})


def test_rounding_matches_excel_and_null_ports_do_not_become_literal_none(templates):
    assert str(exporter._invoice_money("123.445", "Freight")) == "123.45"
    with pytest.raises(ValueError, match="reference number"):
        exporter.generate_pi_invoice_excel(pi(), {"portOfDischarge": None})


def test_export_checks_whole_pi_authorization_before_template_access(monkeypatch):
    monkeypatch.setattr(route, "get_pi_detail", Mock(return_value=pi()))
    denied = Mock(side_effect=HTTPException(403, "Outside authorized country / brand"))
    monkeypatch.setattr(route, "_validate_whole_pi_access", denied)
    generate = Mock()
    monkeypatch.setattr(route, "generate_pi_invoice_excel", generate)
    with pytest.raises(HTTPException) as error:
        route.export_pi_invoice("PI-SE-202609-001", {}, Mock(), Mock())
    assert error.value.status_code == 403
    generate.assert_not_called()


def test_unconfigured_template_returns_actionable_error(tmp_path, monkeypatch):
    monkeypatch.setattr(exporter, "PI_TEMPLATE_ROOT", tmp_path)
    with pytest.raises(ValueError, match="contact|管理员"):
        exporter.pi_invoice_context(pi())


@pytest.mark.parametrize("country,qty_col,total_col,date_cell,ref_cell", [
    ("CH", "F", "L", "J9", "J10"), ("SE", "H", "M", "K8", "K9"),
])
def test_one_template_is_reusable_for_monthly_plans_without_fake_vehicle_units(
    templates, country, qty_col, total_col, date_cell, ref_cell,
):
    line = {**pi(country)["vehicles"][0], "quantity": 54}
    line.pop("freightEur"); line.pop("insuranceEur")
    plan = {"country": country, "month": "2026-10", "lines": [line]}
    before = deepcopy(plan)
    original = (templates / f"{country}_OJ.xlsx").read_bytes()
    options = {"invoiceDate": "2026-10-25", "freightEur": 1700, "insuranceEur": 20}
    output = exporter.generate_monthly_invoice_excel(plan, options)
    values = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    assert values[f"{qty_col}29"].value == 54
    assert values[f"{total_col}29"].value == 632885.94
    assert str(values[date_cell].value).startswith("2026-10-25")
    assert values[ref_cell].value == f"PI-{country}-202610"
    assert plan == before
    plan.update(month="2026-11", lines=[{**line, "quantity": 4, "materialCode": "NEW-MATERIAL", "fobEur": 12000}])
    later = openpyxl.load_workbook(exporter.generate_monthly_invoice_excel(plan, {**options, "invoiceDate": "2026-11-02"}), data_only=True).active
    assert later[f"{qty_col}29"].value == 4 and later[f"{total_col}29"].value == 54880
    assert later[ref_cell].value == f"PI-{country}-202611"
    assert str(later[date_cell].value).startswith("2026-11-02")
    assert (templates / f"{country}_OJ.xlsx").read_bytes() == original


def test_monthly_plan_requires_explicit_costs(templates):
    plan = {"country": "CH", "month": "2026-10", "lines": [{**pi("CH")["vehicles"][0], "quantity": 2}]}
    with pytest.raises(ValueError, match="Freight"):
        exporter.generate_monthly_invoice_excel(plan, {})
    out = exporter.generate_monthly_invoice_excel(plan, {"freightEur": 0, "insuranceEur": 0})
    assert openpyxl.load_workbook(out, data_only=True).active["L29"].value == 20000.22


@pytest.fixture
def model_plan():
    base = pi("CH")["vehicles"][0]
    return {"country": "CH", "month": "2026-10", "lines": [
        {**base, "brand": brand, "modelName": model, "version": trim, "quantity": qty, "materialCode": f"M-{i}"}
        for i, (brand, model, trim, qty) in enumerate([
            ("JAECOO", "MODEL A", "Select", 2), ("JAECOO", "MODEL A", "Exclusive", 3),
            ("JAECOO", "MODEL B", "Select", 5), ("OMODA", "MODEL A", "Select", 1),
        ])]}


def test_monthly_model_groups_use_brand_and_model_not_trim_and_do_not_mutate(templates, model_plan):
    before = deepcopy(model_plan)
    context = exporter.monthly_invoice_context(model_plan)
    assert context["freightModels"] == [
        {"brand": "JAECOO", "modelName": "MODEL A", "unitCount": 5},
        {"brand": "JAECOO", "modelName": "MODEL B", "unitCount": 5},
        {"brand": "OMODA", "modelName": "MODEL A", "unitCount": 1},
    ]
    costs = [{**model, "freightEur": rate} for model, rate in zip(context["freightModels"], (1700, 1750, 0))]
    out = exporter.generate_monthly_invoice_excel(model_plan, {"modelFreightEur": costs, "insuranceEur": 20})
    ws = openpyxl.load_workbook(BytesIO(out.getvalue()), data_only=True).active
    assert [ws[f"I{row}"].value for row in range(28, 32)] == [1700, 1700, 1750, 0]
    assert ws["F32"].value == 11 and ws["L32"].value == 127471.21
    formulas = openpyxl.load_workbook(out).active
    assert formulas["L32"].value == "=SUM(L28:L31)"
    uniform = openpyxl.load_workbook(exporter.generate_monthly_invoice_excel(model_plan,
        {"freightEur": 1700, "insuranceEur": 20}), data_only=True).active
    assert [uniform[f"I{row}"].value for row in range(28, 32)] == [1700] * 4
    assert uniform["L32"].value == 128921.21
    assert model_plan == before


@pytest.mark.parametrize("groups", [None, {}, "1700", [],
    [{"brand": "JAECOO", "modelName": "MODEL A", "freightEur": 1700}],
    [{"brand": "OTHER", "modelName": "MODEL A", "freightEur": 1700}],
    [{"brand": [], "modelName": "MODEL A", "freightEur": 1700}],
    [None],
    [{"brand": "JAECOO", "modelName": "MODEL A", "freightEur": 1}] * 3,
])
def test_monthly_model_freight_never_guesses_missing_extra_duplicate_or_malformed_groups(templates, model_plan, groups):
    with pytest.raises(ValueError, match="freight|Freight|运费"):
        exporter.generate_monthly_invoice_excel(model_plan, {"modelFreightEur": groups, "insuranceEur": 20})


@pytest.mark.parametrize("value", [None, "", -1, float("nan"), float("inf"), 1e12, True])
def test_monthly_model_amounts_reuse_money_validation(templates, model_plan, value):
    groups = [{**model, "freightEur": value} for model in exporter.monthly_invoice_context(model_plan)["freightModels"]]
    with pytest.raises(ValueError, match="Freight|运费"):
        exporter.generate_monthly_invoice_excel(model_plan, {"modelFreightEur": groups, "insuranceEur": 20})


def test_monthly_freight_rejects_both_modes_even_with_a_uniform_fallback(templates, model_plan):
    with pytest.raises(ValueError, match="Choose uniform or model"):
        exporter.generate_monthly_invoice_excel(model_plan, {"modelFreightEur": [], "freightEur": 1700, "insuranceEur": 20})


@pytest.mark.parametrize("output_format", ["template", "details", "data"])
def test_monthly_export_denies_country_before_reading_plan_or_template(monkeypatch, output_format):
    monkeypatch.setattr(route, "_validate_country", Mock(side_effect=HTTPException(403, "Unauthorized country")))
    read = Mock(); monkeypatch.setattr(route, "list_pi_headers", read)
    with pytest.raises(HTTPException) as error:
        route.preview_monthly_invoice("CH", "2026-10", output_format=output_format, session=Mock(), user=Mock())
    assert error.value.status_code == 403
    read.assert_not_called()


@pytest.mark.parametrize("month", ["", "2026-13", "202610", "2026-10-01"])
def test_monthly_export_rejects_missing_or_invalid_month(month):
    with pytest.raises(HTTPException) as error:
        route.preview_monthly_invoice("CH", month, session=Mock(), user=Mock())
    assert error.value.status_code == 400


def test_saved_monthly_summary_preserves_distinct_prices_costs_and_sources(templates):
    first = pi("CH")
    second = pi("CH")
    for index, vehicle in enumerate(first["vehicles"]):
        vehicle.update(piCode="PI-CH-202610-001", carCode=f"A{index}")
    second["lines"][0]["piLineCode"] = "L02"
    for index, vehicle in enumerate(second["vehicles"]):
        vehicle.update(piLineCode="L02", piCode="PI-CH-202610-002", carCode=f"B{index}", fobEur=20000, freightEur=200)
    detail = {"header": {"piCode": "Monthly-Summary-CH-202610", "countryCode": "CH"},
        "lines": first["lines"] + second["lines"], "vehicles": first["vehicles"] + second["vehicles"],
        "sourcePiCodes": ["PI-CH-202610-001", "PI-CH-202610-002"], "summaryExport": True}
    before = deepcopy(detail)
    output = exporter.generate_pi_invoice_excel(detail, {})
    wb = openpyxl.load_workbook(output, data_only=True)
    assert wb.active["F30"].value == 4
    assert wb.active["G28"].value == 10000.11 and wb.active["G29"].value == 20000
    assert wb.active["I28"].value == 123.45 and wb.active["I29"].value == 200
    assert wb.active["H30"].value == 60000.22
    assert "Monthly PI summary" in wb["PI sources"]["A1"].value
    assert {wb["PI sources"].cell(row, 1).value for row in range(3, 7)} == set(detail["sourcePiCodes"])
    assert detail == before
    grouped = exporter.generate_pi_invoice_excel(detail, {"modelFreightEur": [{"brand": "OMODA", "modelName": "MODEL BEV", "freightEur": 0}]})
    assert openpyxl.load_workbook(grouped, data_only=True).active["I28"].value == 0


def test_formal_monthly_export_requires_saved_source_and_current_preview(monkeypatch):
    monkeypatch.setattr(route, "_monthly_invoice_source", Mock(return_value=(None, "CH", {"JAECOO"}, [])))
    generate = Mock(); monkeypatch.setattr(route, "generate_pi_invoice_excel", generate)
    with pytest.raises(HTTPException) as error:
        route.export_monthly_invoice("CH", "2026-10", {}, session=Mock(), user=Mock())
    assert error.value.status_code == 409 and error.value.detail["code"] == "no_saved_pi"
    detail = {**pi("CH"), "sourceRevision": "new"}
    monkeypatch.setattr(route, "_monthly_invoice_source", Mock(return_value=(detail, "CH", {"OMODA"}, [])))
    with pytest.raises(HTTPException) as error:
        route.export_monthly_invoice("CH", "2026-10", {"sourceRevision": "old"}, session=Mock(), user=Mock())
    assert error.value.detail["code"] == "source_changed"
    generate.assert_not_called()


@pytest.mark.parametrize("output_format,first", [("details", 5), ("data", 2)])
def test_template_free_formats_use_same_saved_snapshots_and_live_formulas(tmp_path, monkeypatch, output_format, first):
    monkeypatch.setattr(exporter, "PI_TEMPLATE_ROOT", tmp_path / "no-templates")
    detail = pi("BG")
    detail["vehicles"][0].update(piCode=detail["header"]["piCode"], carCode="C1", modelName="=1+1")
    detail["vehicles"][1].update(piCode=detail["header"]["piCode"], carCode="C2", fobEur=20000, freightEur=200)
    before = deepcopy(detail)
    context = exporter.pi_invoice_context(detail, output_format=output_format)
    assert context["unitCount"] == 2 and context["buyerName"] == "" and context["bankDetails"] == ""
    output = exporter.generate_pi_invoice_excel(detail, {"outputFormat": output_format})
    wb = openpyxl.load_workbook(output)
    values = openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=True).active
    ws = wb.active
    assert values[f"H{first + 2}"].value == 2
    assert values[f"J{first + 2}"].value == 30000.11
    assert values[f"N{first + 2}"].value == 30337.12
    assert ws[f"J{first}"].value == f"=ROUND(H{first}*I{first},2)"
    assert values[f"I{first}"].value == 10000.11 and values[f"I{first + 1}"].value == 20000
    assert ws[f"B{first}"].value == "'=1+1" and ws[f"B{first}"].data_type == "s"
    assert ws.freeze_panes == f"H{first}" and ws.page_setup.fitToHeight == 0
    assert ws.print_area == f"'{ws.title}'!$A$1:$N${first + 2}"
    if output_format == "data":
        assert not ws.merged_cells.ranges and ws.auto_filter.ref == "A1:N3"
        assert {wb["PI sources"].cell(row, 1).value for row in (3, 4)} == {detail["header"]["piCode"]}
    else:
        assert ws["A1"].value == "9-Description and Quantity: Brand New SUV Motor Vehicles"
        assert "Description of Goods" == ws["B3"].value
    assert detail == before


@pytest.mark.parametrize("output_format", ["template", "details", "data"])
def test_all_three_formats_share_amounts_and_explicit_model_costs(templates, output_format):
    detail = pi("CH")
    output = exporter.generate_pi_invoice_excel(detail, {"outputFormat": output_format,
        "modelFreightEur": [{"brand": "OMODA", "modelName": "MODEL BEV", "freightEur": 0}],
        "insuranceEur": 0, "handlingEur": 5})
    ws = openpyxl.load_workbook(output, data_only=True).active
    total, quantity, amount = (29, "F", "L") if output_format == "template" else (6 if output_format == "details" else 3, "H", "N")
    assert ws[f"{quantity}{total}"].value == 2 and ws[f"{amount}{total}"].value == 20010.22
    detail["vehicles"][0]["fobEur"] = None
    with pytest.raises(ValueError, match="FOB"):
        exporter.generate_pi_invoice_excel(detail, {"outputFormat": output_format})
    with pytest.raises(ValueError):
        exporter.generate_pi_invoice_excel(pi("CH"), {"outputFormat": output_format, "insuranceEur": -1})


@pytest.mark.parametrize("output_format", ["details", "data"])
def test_template_free_routes_do_not_read_template_files_or_bypass_saved_pi(monkeypatch, output_format):
    detail = {**pi("BG"), "sourceRevision": "r1", "sourcePiCodes": ["PI-BG-202609-001"]}
    monkeypatch.setattr(route, "_monthly_invoice_source", Mock(return_value=(detail, "BG", {"OMODA"}, [])))
    template_read = Mock(side_effect=AssertionError("No template dependency allowed"))
    monkeypatch.setattr(route, "_template_catalog", template_read)
    monkeypatch.setattr(route, "_scoped_template", template_read)
    preview = route.preview_monthly_invoice("BG", "2026-09", output_format=output_format, session=Mock(), user=Mock())
    assert preview["context"]["sourceRevision"] == "r1" and preview["catalog"] is None
    response = route.export_monthly_invoice("BG", "2026-09", {"outputFormat": output_format, "sourceRevision": "r1"}, session=Mock(), user=Mock())
    assert response.media_type == route.XLSX_MEDIA_TYPE and "PI-BG-202609-001" in response.headers["content-disposition"]
    monkeypatch.setattr(route, "_monthly_invoice_source", Mock(return_value=(None, "BG", {"OMODA"}, [])))
    assert route.preview_monthly_invoice("BG", "2026-09", output_format=output_format, session=Mock(), user=Mock())["issue"]["code"] == "no_saved_pi"
    with pytest.raises(HTTPException) as error:
        route.export_monthly_invoice("BG", "2026-09", {"outputFormat": output_format}, session=Mock(), user=Mock())
    assert error.value.status_code == 409 and error.value.detail["code"] == "no_saved_pi"
    template_read.assert_not_called()


@pytest.mark.parametrize("output_format,first", [("details", 5), ("data", 2)])
def test_template_free_http_download_and_role_gate(monkeypatch, output_format, first):
    detail = {**pi("BG"), "sourceRevision": "r1", "sourcePiCodes": ["PI-BG-202609-001"]}
    monkeypatch.setattr(route, "_monthly_invoice_source", Mock(return_value=(detail, "BG", {"OMODA"}, [])))
    monkeypatch.setattr(route, "_scoped_template", Mock(side_effect=AssertionError("Must not load a template")))
    app = FastAPI(); app.include_router(route.router)
    user = UserContext("order_filler", "alice")
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[route.get_db_session] = lambda: Mock()
    path = f"/order-genius/vehicle-allocation/monthly-invoice?country=BG&month=2026-09&output_format={output_format}"
    with TestClient(app) as client:
        preview = client.get(path)
        assert preview.status_code == 200 and preview.json()["context"]["unitCount"] == 2
        response = client.post(path, json={"outputFormat": output_format, "sourceRevision": "r1"})
        assert response.status_code == 200 and response.headers["content-type"] == route.XLSX_MEDIA_TYPE
        ws = openpyxl.load_workbook(BytesIO(response.content), data_only=True).active
        assert ws[f"N{first + 1}"].value == 20260.68
        assert client.post(path, json={"outputFormat": output_format, "sourceRevision": "old"}).status_code == 409
        user = UserContext("viewer", "alice")
        assert client.get(path).status_code == 403 and client.post(path, json={}).status_code == 403


@pytest.mark.parametrize("output_format", [None, [], {}, "unknown"])
def test_invalid_export_formats_are_rejected_without_reading_templates(output_format):
    with pytest.raises(ValueError, match="Choose template"):
        exporter.generate_pi_invoice_excel(pi("BG"), {"outputFormat": output_format})
