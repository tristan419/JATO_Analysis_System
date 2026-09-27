"""Human-readable BOM Admin workbook export.

The workbook mirrors the external material-code reference layout: one model per
sheet, configuration/BOM groupings, colour rows, and country FOB columns.  It
also carries explicit machine-readable tier and price components so a later
preview import never has to infer pricing from colour names or cell fills.
"""

from __future__ import annotations

import io
import hashlib
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from typing import Iterable
from xml.etree import ElementTree

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.infra.order_genius_repository import COUNTRY_NAMES_BY_CODE


BOM_ADMIN_EXPORT_SCHEMA_VERSION = "bom-admin-material-master-v1"

_HEADER_FONT = Font(name="Calibri", size=14, bold=True, color="FF000000")
_SURCHARGE_FILL = PatternFill("solid", fgColor="FFE1EAFF")
_CONFLICT_FILL = PatternFill("solid", fgColor="FFFDE9E7")
_THIN_GREY = Side(style="thin", color="FF7F7F7F")
_BORDER = Border(
    left=_THIN_GREY,
    right=_THIN_GREY,
    top=_THIN_GREY,
    bottom=_THIN_GREY,
)
_PRICE_FORMAT = '#,##0 "€";[Red](#,##0 "€");-'

_FIXED_HEADERS = (
    "No.",
    "Code Name",
    "Full Name",
    "Configuration",
    "BOM",
    "Material Code",
    "Exterior Color",
    "Colour Code",
    "Interior Color",
    "Colour Tier",
)

_COLOUR_TIERS = frozenset({"single", "dual", "special"})


def _clean(value: object) -> str:
    return str(value or "").strip()


def _safe_sheet_name(raw_name: str, used_names: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", _clean(raw_name)) or "MODEL"
    base = re.sub(r"\s+", " ", base).strip()[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used_names:
        marker = f" {suffix}"
        candidate = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _normalise_countries(countries: Iterable[str], country_code: str | None) -> list[str]:
    available = sorted({_clean(code).upper() for code in countries if _clean(code)})
    if country_code:
        selected = _clean(country_code).upper()
        if selected not in available:
            raise ValueError(f"Country {selected} has no active BOM FOB data")
        return [selected]
    return available


def _country_header(country_code: str) -> str:
    if country_code == "CZ":
        return "CZ FOB"
    return f"{COUNTRY_NAMES_BY_CODE.get(country_code, country_code)} FOB"


def _group_sort_key(item: dict) -> tuple[str, str, str, int, str]:
    tier = _clean(item.get("colourTier")).lower()
    return (
        _clean(item.get("version")).casefold(),
        _clean(item.get("bomTemplate")).casefold(),
        _clean(item.get("interiorColorName")).casefold(),
        {"single": 0, "dual": 1, "special": 2}.get(tier, 3),
        _clean(item.get("colourCode")).casefold(),
    )


def _explicit_colour_tier(item: dict) -> str | None:
    tier = _clean(item.get("colourTier")).lower()
    return tier if tier in _COLOUR_TIERS else None


def _country_price_parts(
    item: dict,
    country_code: str,
) -> tuple[object, object, object, str, str]:
    """Return persisted price components without reconstructing business data."""
    fob = (item.get("fobByCountry") or {}).get(country_code) or {}
    if fob.get("status") == "conflict":
        return None, None, None, "conflict", json.dumps(
            fob.get("records") or [],
            ensure_ascii=False,
            sort_keys=True,
        )

    final = fob.get("finalFobEur")
    base = fob.get("baseFobEur")
    surcharge = fob.get("colourSurchargeEur")
    tier = _explicit_colour_tier(item)
    if tier is None:
        status = "missing_tier"
    elif tier == "single" and surcharge is None:
        surcharge = 0
        status = "ready" if base is not None and final is not None else "missing_price"
    elif base is None:
        status = "missing_single_base"
    elif surcharge is None:
        status = "missing_surcharge"
    elif final is None:
        status = "missing_final_fob"
    else:
        status = "ready"

    return base, surcharge, final, status, ""


def _merge_group_columns(ws, start_row: int, end_row: int) -> None:
    if end_row <= start_row:
        return
    for column in (4, 5, 9):
        ws.merge_cells(
            start_row=start_row,
            start_column=column,
            end_row=end_row,
            end_column=column,
        )
        ws.cell(start_row, column).alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )


def _write_model_sheet(ws, items: list[dict], countries: list[str]) -> None:
    fixed_count = len(_FIXED_HEADERS)
    for column, label in enumerate(_FIXED_HEADERS, start=1):
        cell = ws.cell(1, column, label)
        ws.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER

    for country_index, country_code in enumerate(countries):
        start_col = fixed_count + country_index + 1
        title = ws.cell(1, start_col, _country_header(country_code))
        ws.merge_cells(start_row=1, start_column=start_col, end_row=2, end_column=start_col)
        title.font = _HEADER_FONT
        title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        title.border = _BORDER

    grouped: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for item in sorted(items, key=_group_sort_key):
        grouped[
            (
                _clean(item.get("version")),
                _clean(item.get("bomTemplate")),
                _clean(item.get("interiorColorName")),
            )
        ].append(item)

    row = 3
    model_start = row
    group_entries = list(grouped.items())
    for group_index, ((version, bom_template, interior), group_items) in enumerate(group_entries):
        group_start = row
        for index, item in enumerate(group_items, start=1):
            tier = _explicit_colour_tier(item)
            exterior_name = _clean(item.get("colour"))
            colour_code = _clean(item.get("colourCode"))
            values = (
                index,
                _clean(item.get("modelCode")),
                _clean(item.get("modelName")),
                version,
                bom_template,
                _clean(item.get("materialCode")),
                f"{exterior_name} ({colour_code})" if colour_code else exterior_name,
                colour_code,
                interior,
                tier.title() if tier else "Missing tier",
            )
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row, column, value)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = _BORDER
                if tier in {"dual", "special"}:
                    cell.fill = _SURCHARGE_FILL
                if column in {5, 6, 8}:
                    cell.number_format = "@"
                if column == 10 and not tier:
                    cell.fill = _CONFLICT_FILL
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FF9C0006")

            for country_index, country_code in enumerate(countries):
                column = fixed_count + country_index + 1
                _, _, final, price_status, _ = _country_price_parts(item, country_code)
                value = "REVIEW" if price_status == "conflict" else final
                cell = ws.cell(row, column, value)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = _BORDER
                if isinstance(value, (int, float)):
                    cell.number_format = _PRICE_FORMAT
                if price_status == "conflict":
                    cell.fill = _CONFLICT_FILL
                elif tier in {"dual", "special"}:
                    cell.fill = _SURCHARGE_FILL
            row += 1

        _merge_group_columns(ws, group_start, row - 1)
        if group_index < len(group_entries) - 1:
            row += 1

    model_end = row - 1
    if model_end >= model_start:
        while model_end > model_start and all(
            ws.cell(model_end, column).value is None
            for column in range(1, fixed_count + len(countries) + 1)
        ):
            model_end -= 1
        for column in (2, 3):
            ws.merge_cells(
                start_row=model_start,
                start_column=column,
                end_row=model_end,
                end_column=column,
            )
            ws.cell(model_start, column).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    widths = (8, 16, 22, 20, 30, 24, 42, 12, 20, 14)
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for column in range(fixed_count + 1, fixed_count + len(countries) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 14

    ws.row_dimensions[1].height = 27
    ws.row_dimensions[2].height = 10
    ws.freeze_panes = f"{get_column_letter(fixed_count + 1)}3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"
    ws.sheet_properties.outlinePr.summaryBelow = True


def _price_fingerprint(item: dict, country_code: str, fob: dict) -> str:
    payload = {
        "materialCode": _clean(item.get("materialCode")),
        "bomTemplate": _clean(item.get("bomTemplate")),
        "colourTier": _explicit_colour_tier(item),
        "rowVersion": item.get("rowVersion"),
        "countryCode": country_code,
        "fob": fob,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _write_data_sheet(wb, items: list[dict], countries: list[str]) -> None:
    ws = wb.create_sheet("_BOM_ADMIN_DATA")
    headers = (
        "Schema Version",
        "Brand",
        "Model Code",
        "Full Name",
        "Configuration",
        "BOM Template",
        "Material Code",
        "Exterior Color",
        "Colour Code",
        "Colour HEX",
        "Interior Color",
        "Colour Tier",
        "Row Version",
        "Country Code",
        "Single Base",
        "Surcharge",
        "Final FOB",
        "Price Status",
        "Conflict Detail",
        "Fingerprint",
    )
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(1, column, header)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.border = _BORDER

    row = 2
    for item in sorted(items, key=lambda entry: (
        _clean(entry.get("modelName")).casefold(),
        _group_sort_key(entry),
    )):
        tier = _explicit_colour_tier(item)
        for country_code in countries:
            fob = (item.get("fobByCountry") or {}).get(country_code) or {}
            base, surcharge, final, status, conflict_detail = _country_price_parts(
                item,
                country_code,
            )
            values = (
                BOM_ADMIN_EXPORT_SCHEMA_VERSION,
                _clean(item.get("brand")),
                _clean(item.get("modelCode")),
                _clean(item.get("modelName")),
                _clean(item.get("version")),
                _clean(item.get("bomTemplate")),
                _clean(item.get("materialCode")),
                _clean(item.get("colour")),
                _clean(item.get("colourCode")),
                _clean(item.get("colourHex")),
                _clean(item.get("interiorColorName")),
                tier or "",
                item.get("rowVersion"),
                country_code,
                base,
                surcharge,
                final,
                status,
                conflict_detail,
                _price_fingerprint(item, country_code, fob),
            )
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row, column, value)
                cell.border = _BORDER
                if column in {6, 7, 9, 10, 20}:
                    cell.number_format = "@"
                elif column in {15, 16, 17} and isinstance(value, (int, float)):
                    cell.number_format = _PRICE_FORMAT
            row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:T{max(row - 1, 1)}"
    ws.sheet_state = "hidden"


def _write_schema_sheet(wb, countries: list[str], item_count: int) -> None:
    ws = wb.create_sheet("_BOM_ADMIN_SCHEMA")
    rows = (
        ("schemaVersion", BOM_ADMIN_EXPORT_SCHEMA_VERSION),
        ("exportedAtUtc", datetime.now(timezone.utc).isoformat()),
        ("source", "BOM Admin / Material Master"),
        ("priceModel", "template-country Single base + saved colour tier surcharge"),
        ("tierAuthority", "BOM Admin saved Colour Tier; never inferred from name/code/swatch/fill"),
        ("countryScope", ",".join(countries)),
        ("itemCount", item_count),
    )
    for row_index, (key, value) in enumerate(rows, start=1):
        ws.cell(row_index, 1, key)
        ws.cell(row_index, 2, value)
    ws.sheet_state = "hidden"


def _normalise_font_xml_order(workbook_bytes: io.BytesIO) -> io.BytesIO:
    """Keep generated XLSX font children in the strict OpenXML schema order."""
    source = workbook_bytes.getvalue()
    output = io.BytesIO()
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    order = {
        "b": 0,
        "i": 1,
        "strike": 2,
        "outline": 3,
        "shadow": 4,
        "condense": 5,
        "extend": 6,
        "sz": 7,
        "color": 8,
        "name": 9,
        "family": 10,
        "scheme": 11,
        "charset": 12,
    }
    ElementTree.register_namespace("", namespace)
    with zipfile.ZipFile(io.BytesIO(source), "r") as source_zip:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target_zip:
            for entry in source_zip.infolist():
                data = source_zip.read(entry.filename)
                if entry.filename == "xl/styles.xml":
                    root = ElementTree.fromstring(data)
                    fonts = root.find(f"{{{namespace}}}fonts")
                    if fonts is not None:
                        for font in fonts:
                            children = list(font)
                            children.sort(
                                key=lambda child: order.get(
                                    child.tag.rsplit("}", 1)[-1],
                                    len(order),
                                )
                            )
                            font[:] = children
                    data = ElementTree.tostring(
                        root,
                        encoding="utf-8",
                        xml_declaration=True,
                    )
                target_zip.writestr(entry, data)
    output.seek(0)
    return output


def build_bom_admin_workbook(
    items: list[dict],
    countries: Iterable[str],
    *,
    country_code: str | None = None,
) -> io.BytesIO:
    """Build the BOM Admin material workbook for all countries or one country."""
    selected_countries = _normalise_countries(countries, country_code)
    if not selected_countries:
        raise ValueError("No active BOM FOB countries are available for export")

    model_groups: dict[str, list[dict]] = defaultdict(list)
    for item in items:
        model_name = _clean(item.get("modelName")) or "UNNAMED MODEL"
        model_groups[model_name].append(item)
    if not model_groups:
        raise ValueError("No BOM Admin material rows are available for export")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for model_name in sorted(model_groups, key=str.casefold):
        ws = wb.create_sheet(_safe_sheet_name(model_name, used_names))
        _write_model_sheet(ws, model_groups[model_name], selected_countries)
    _write_data_sheet(wb, items, selected_countries)
    _write_schema_sheet(wb, selected_countries, len(items))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return _normalise_font_xml_order(output)
