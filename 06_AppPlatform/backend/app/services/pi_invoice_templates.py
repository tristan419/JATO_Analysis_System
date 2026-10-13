"""Personal PI templates and explicit country/brand sharing; no order writes."""
from __future__ import annotations

import hashlib
import io
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile, BadZipFile

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from upload_toolkit.job_engine import load_job_state, persist_job_state

MAX_TEMPLATE_BYTES = 8 * 1024 * 1024
FIELDS = ("sellerName", "sellerAddress", "buyerName", "buyerAddress", "buyerEmail", "bankDetails",
          "priceTerm", "paymentTerm", "invoiceDate", "referenceNo", "portOfShipment", "portOfDischarge")
REQUIRED_FIELDS = {"sellerName", "buyerName", "buyerAddress", "bankDetails", "priceTerm", "paymentTerm",
                   "invoiceDate", "referenceNo", "portOfShipment", "portOfDischarge"}
COLUMNS = ("number", "modelName", "version", "powertrain", "exteriorColorName", "interiorColorName", "materialCode",
           "quantity", "fobEur", "fobAmount", "freightEur", "insuranceEur", "handlingEur", "total")
ALIASES = {
    "number": ("no", "number"), "modelName": ("model", "车型"), "version": ("version", "版型", "配置"),
    "powertrain": ("displacementbatterycapacity", "powertrain", "动力"),
    "exteriorColorName": ("color", "colour", "exteriorcolor", "exteriorcolour", "外饰"),
    "interiorColorName": ("interior", "interiorcolor", "内饰"), "materialCode": ("material", "materialcode", "物料号"),
    "quantity": ("quantity", "qty", "数量"), "fobEur": ("fob", "unitpricefob", "unitfob", "fobperunit"),
    "fobAmount": ("fobamount", "totalamountfob", "amountfob", "totalfob"), "freightEur": ("freight", "运费"),
    "insuranceEur": ("insurance", "保费"), "handlingEur": ("handlingcharges", "handingcharges", "handling", "操作费"),
    "total": ("totalamountcif", "totalamountcfr", "cifamount", "cframount", "totalamount", "totalprice", "totalcif", "totalcfr", "合计金额"),
}


def root() -> Path:
    from app.services.order_genius_vehicle_exporter import PI_TEMPLATE_ROOT
    return PI_TEMPLATE_ROOT


def _normal(value: object) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", str(value or "").lower().replace("eur", ""))


def workbook_bytes(content: bytes):
    if not content or len(content) > MAX_TEMPLATE_BYTES:
        raise ValueError("Template must be an XLSX under 8 MiB / 模板须为8MiB以内的XLSX")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("Template expands beyond 40 MiB / 模板解压后过大")
            if any("vbaProject" in name or "externalLinks/" in name for name in names):
                raise ValueError("Macros/external workbooks are unsupported / 不支持宏或外部工作簿链接")
        workbook = openpyxl.load_workbook(io.BytesIO(content), keep_links=False)
        if len(workbook.sheetnames) != 1:
            workbook.close()
            raise ValueError("First release supports one invoice worksheet / 本轮支持单工作表PI模板，请先单独保存主发票页")
        return workbook
    except (BadZipFile, OSError, KeyError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError("Invalid XLSX template / XLSX模板无法读取") from exc


def _check_sheet(sheet) -> None:
    if sheet.max_column > 40 or sheet.max_row > 2000:
        raise ValueError("Use one PI table within 40 columns / 2000 rows / 请使用40列、2000行以内的单一PI明细表")
    if any(re.search(r"\b(?:USD|CHF|SEK|GBP)\b", str(cell.value or ""), re.I)
               for row in sheet.iter_rows(max_row=min(100, sheet.max_row)) for cell in row if cell.data_type != "f"):
        raise ValueError("Saved PI prices are EUR; upload an EUR template / 已保存PI价格为EUR，请上传EUR模板；不自动换汇")


def analyse(content: bytes, sheet_name: str = "") -> dict:
    workbook = workbook_bytes(content)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        _check_sheet(sheet)
        best, header_row, header_end = {}, 0, 0
        for row in range(1, min(sheet.max_row, 100) + 1):
            found = {}
            for column in range(1, sheet.max_column + 1):
                # Two-row headings may have the label at either side of a vertical merge.
                labels = {_normal(sheet.cell(r, column).value) for r in (max(1, row - 1), row)}
                for field, aliases in ALIASES.items():
                    if labels.intersection(aliases) and field not in found:
                        found[field] = get_column_letter(column)
            if {"modelName", "quantity", "fobEur"} <= found.keys() and len(found) > len(best):
                current_labels = {_normal(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
                one_row = all(current_labels.intersection(ALIASES[key]) for key in ("modelName", "quantity", "fobEur"))
                best, header_row, header_end = found, row if one_row else max(1, row - 1), row
        start = header_end + 1 if header_end else 0
        total = next((row for row in range(max(1, start), sheet.max_row + 1)
                      if _normal(sheet.cell(row, 1).value) in {"total", "合计", "总计"}), 0)
        cells = [{"cell": cell.coordinate, "value": str(cell.value)[:3000]} for row in sheet.iter_rows(max_row=min(200, sheet.max_row), max_col=sheet.max_column)
                 for cell in row if cell.value is not None and cell.data_type != "f" and not (start and total and start <= cell.row <= total)]
        fields = {key: {"cell": "", "value": "", "prefix": ""} for key in FIELDS}
        # Label-based candidates only. Every candidate is explicitly reviewed before saving.
        patterns = {"buyerName": r"^(?:m/s|buyer)\s*[:：]", "buyerAddress": r"^address\s*[:：]", "buyerEmail": r"^email\s*[:：]",
                    "invoiceDate": r"^(?:date|invoice date)\s*[:：]?", "referenceNo": r"^(?:ref\.?\s*no|invoice no)\s*[:：]?",
                    "portOfShipment": r"(?:port of shipment|loading port)\s*[:：]", "portOfDischarge": r"(?:port of discharge|discharge port)\s*[:：]"}
        for key, pattern in patterns.items():
            matches = [cell for cell in cells if re.search(pattern, cell["value"], re.I)]
            if len(matches) == 1:
                value = matches[0]["value"]
                prefix = value[:re.search(r"[:：]", value).end()] + " " if re.search(r"[:：]", value) else ""
                fields[key] = {"cell": matches[0]["cell"], "value": value[len(prefix.rstrip()):].strip(), "prefix": prefix}
        return {"sha256": hashlib.sha256(content).hexdigest(), "sheets": workbook.sheetnames,
                "sheetName": sheet.title, "headerRow": header_row, "dataStart": start, "totalRow": total,
                "columns": best, "fields": fields, "sourceCells": cells,
                "charges": [], "currency": "EUR"}
    finally:
        workbook.close()


def validate_config(content: bytes, config: dict) -> dict:
    if not isinstance(config, dict):
        raise ValueError("Invalid template mapping / 模板映射无效")
    workbook = workbook_bytes(content)
    try:
        sheet = workbook[config["sheetName"]]
        _check_sheet(sheet)
        header, start, total = (config.get(key) for key in ("headerRow", "dataStart", "totalRow"))
        if any(type(value) is not int for value in (header, start, total)) or not (1 <= header < start < total <= sheet.max_row + 1):
            raise ValueError("Check header, first detail and total rows / 请核对表头、明细起始行和总计行")
        columns = config.get("columns", {})
        if not isinstance(columns, dict) or not {"modelName", "quantity", "fobEur", "total"} <= columns.keys() or set(columns) - set(COLUMNS):
            raise ValueError("Map Model, Quantity, FOB and Total / 必须映射车型、数量、FOB和合计列")
        indexes = [column_index_from_string(str(value)) for value in columns.values()]
        if len(indexes) != len(set(indexes)) or not all(1 <= value <= min(40, sheet.max_column) for value in indexes):
            raise ValueError("Columns overlap or are outside the table / 列重复或超出模板范围")
        charges = config.get("charges")
        if not isinstance(charges, list) or not all(isinstance(value, str) for value in charges) or len(charges) != len(set(charges)) or set(charges) - {"freightEur", "insuranceEur", "handlingEur"} or not set(charges) <= columns.keys():
            raise ValueError("Confirm the charges included in totals / 请确认哪些费用计入合计")
        if config.get("currency") != "EUR":
            raise ValueError("Only saved EUR PI values are supported / 本轮仅支持已保存PI的EUR金额")
        fields = config.get("fields", {})
        if not isinstance(fields, dict) or set(fields) - set(FIELDS):
            raise ValueError("Invalid template fields / 模板字段无效")
        addresses, clean_fields = set(), {}
        max_column, max_row = sheet.max_column, sheet.max_row
        for key in FIELDS:
            field = fields.get(key, {})
            if not isinstance(field, dict):
                raise ValueError("Invalid field mapping / 字段映射无效")
            address = str(field.get("cell") or "").upper().strip()
            if key in REQUIRED_FIELDS and (not address or (key not in {"invoiceDate", "referenceNo", "portOfShipment", "portOfDischarge"} and not str(field.get("value") or "").strip())):
                raise ValueError(f"Confirm {key} and its cell / 请确认{key}及单元格")
            if not address:
                continue
            if not re.fullmatch(r"[A-Z]{1,2}[1-9][0-9]{0,3}", address) or address in addresses:
                raise ValueError("Field cells are invalid or overlap / 字段单元格无效或重复")
            cell = sheet[address]
            if start <= cell.row <= total or cell.column > max_column or cell.row > max_row:
                raise ValueError("Header fields must be outside invoice details / 固定字段不能覆盖明细或总计")
            if any(address in merged and address != merged.start_cell.coordinate for merged in sheet.merged_cells.ranges):
                raise ValueError("Use the top-left cell of a merged field / 合并字段请用左上角单元格")
            if any(len(str(field.get(part) or "")) > 4000 for part in ("value", "prefix")):
                raise ValueError("Template field is too long / 模板字段过长")
            addresses.add(address)
            clean_fields[key] = {"cell": address, "value": str(field.get("value") or ""), "prefix": str(field.get("prefix") or "")}
        for merged in sheet.merged_cells.ranges:
            if merged.min_row < start <= merged.max_row or merged.min_row <= total < merged.max_row:
                raise ValueError("Merged cell crosses invoice boundaries / 合并单元格跨越明细边界，请先调整")
            if merged.min_row == total and any(merged.min_col <= column <= merged.max_col and column != merged.min_col for field, column in zip(columns, indexes) if field in {"quantity", "total", "fobAmount"}):
                raise ValueError("Total cell is inside a merge / 总计映射请用合并单元格左上角或取消该合并")
        # First release regenerates the invoice table; unsupported calculations elsewhere cannot survive unnoticed.
        for row in sheet:
            for cell in row:
                if cell.data_type == "f" and not (start <= cell.row <= total) and cell.coordinate not in addresses:
                    raise ValueError(f"Unsupported formula outside invoice table: {cell.coordinate} / 明细外公式需要先调整")
        return {"sheetName": sheet.title, "headerRow": header, "dataStart": start, "totalRow": total,
                "columns": {key: str(value).upper() for key, value in columns.items()}, "fields": clean_fields, "charges": charges, "currency": "EUR"}
    except (KeyError, TypeError, IndexError) as exc:
        raise ValueError("Template mapping is incomplete / 模板映射不完整") from exc
    finally:
        workbook.close()


def _path(template_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", template_id):
        raise ValueError("Invalid template identity / 模板编号无效")
    return root() / "personal" / template_id


def records() -> list[dict]:
    directory = root() / "personal"
    return [load_job_state(path) for path in sorted(directory.glob("*/template.json"))] if directory.exists() else []


def visible(record: dict, actor: str, admin: bool, country: str, brands: set[str], owner: str) -> bool:
    return (record["country"] == country and brands <= set(record["brands"])
            and (admin or record["owner"] == actor or record.get("shared", False))
            and (owner == "*" or record["owner"] == owner or record.get("shared", False)))


def builtin(country: str, brands: set[str]) -> dict | None:
    if country not in {"CH", "SE"} or not brands or not brands <= {"OMODA", "JAECOO"}:
        return None
    profiles = root() / "profiles.json"
    if not profiles.is_file() or not (root() / f"{country}_OJ.xlsx").is_file():
        return None
    profile = load_job_state(profiles).get(country)
    return ({"id": f"{country}_OJ", "name": profile["name"], "owner": "System", "country": country,
             "brands": ["OMODA", "JAECOO"], "shared": True, "system": True} if profile else None)


def catalog(actor: str, admin: bool, country: str, brands: set[str], owner: str = "") -> dict:
    owner = owner or actor
    if owner != actor and not admin:
        raise PermissionError("Only admin can manage another owner / 仅管理员可管理其他账号")
    items = [record for record in records() if visible(record, actor, admin, country, brands, owner)]
    standard = builtin(country, brands)
    if standard:
        items.append(standard)
    order_file = root() / "orders" / (hashlib.sha256(f"{owner}|{country}|{','.join(sorted(brands))}".encode()).hexdigest() + ".json")
    order = load_job_state(order_file).get("ids", []) if order_file.exists() else []
    items.sort(key=lambda item: (order.index(item["id"]) if item["id"] in order else len(order), item.get("system", False), item["owner"] != owner, item.get("createdAt", ""), item["id"]))
    return {"items": items, "owner": owner, "owners": sorted({record["owner"] for record in records()}) if admin else [actor]}


def load_template(template_id: str, actor: str, admin: bool, country: str, brands: set[str]) -> dict:
    record = load_job_state(_path(template_id) / "template.json")
    if not visible(record, actor, admin, country, brands, "*"):
        raise PermissionError("Template is outside your scope / 模板不在当前授权范围")
    content = (_path(template_id) / "source.xlsx").read_bytes()
    if hashlib.sha256(content).hexdigest() != record["sha256"]:
        raise ValueError("Template file changed or is damaged / 模板文件发生变化或损坏")
    return {**record, "path": _path(template_id) / "source.xlsx"}


def save(content: bytes, payload: dict, actor: str, admin: bool) -> dict:
    template_id = str(payload.get("id") or "")
    existing = load_job_state(_path(template_id) / "template.json") if template_id else None
    if existing and existing["owner"] != actor and not admin:
        raise PermissionError("Only the owner or admin can edit this source / 只有所有者或管理员可改源模板")
    digest = hashlib.sha256(content).hexdigest()
    if payload.get("sha256") != digest:
        raise ValueError("Re-parse the selected file before saving / 文件已变化，请重新解析后保存")
    config = validate_config(content, payload["config"])
    name = str(payload.get("name") or "").strip()
    if not name or len(name) > 100 or type(payload.get("shared", False)) is not bool:
        raise ValueError("Check template name and sharing / 请核对模板名称及共享设置")
    template_id = template_id or uuid4().hex
    record = {"id": template_id, "owner": existing["owner"] if existing else actor, "name": name,
              "country": payload["country"], "brands": payload["brands"], "shared": payload.get("shared", False),
              "sha256": digest, "config": config, "createdAt": existing["createdAt"] if existing else datetime.now(timezone.utc).isoformat()}
    directory = _path(template_id)
    directory.mkdir(parents=True, exist_ok=True)
    # Existing edits keep the same verified file; only explicit new imports create a source.
    if not existing:
        (directory / "source.xlsx").write_bytes(content)
    elif digest != existing["sha256"]:
        raise ValueError("Import a changed file as a new template / 文件改版请作为新模板导入")
    persist_job_state(directory / "template.json", record)
    return record


def reorder(actor: str, admin: bool, country: str, brands: set[str], ids: list[str], owner: str = "") -> None:
    owner = owner or actor
    listing = catalog(actor, admin, country, brands, owner)
    if owner == "*" or not isinstance(ids, list) or not all(isinstance(value, str) for value in ids) or len(ids) != len(set(ids)) or set(ids) != {record["id"] for record in listing["items"]}:
        raise ValueError("Reload the complete template list before reordering / 排序前请重读完整模板列表")
    path = root() / "orders" / (hashlib.sha256(f"{owner}|{country}|{','.join(sorted(brands))}".encode()).hexdigest() + ".json")
    persist_job_state(path, {"ids": ids})


def remove(template_id: str, actor: str, admin: bool, country: str, brands: set[str]) -> None:
    record = load_job_state(_path(template_id) / "template.json")
    if not visible(record, actor, admin, country, brands, "*") or (record["owner"] != actor and not admin):
        raise PermissionError("Only the owner or admin can delete this source / 只有所有者或管理员可删除源模板")
    # Exact UUID directory, never a caller-supplied path or the templates root.
    shutil.rmtree(_path(template_id))
